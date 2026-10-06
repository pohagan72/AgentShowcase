"""Build Synzo-MCP-Server-Specification.docx from the Harvey example template.

Strategy: open the Harvey example .docx so we inherit its styles (Title, Heading 1/2/3,
normal, table style), wipe the body, and append the Synzo content using those same
styles. Keeps the look-and-feel the Harvey evaluators expect while replacing every
ACME-specific statement with verified Synzo facts.
"""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


TEMPLATE = Path(r"C:\Users\651802\Downloads\[Example] Ideal MCP Server Specification - Harvey AI (1).docx")
OUTPUT = Path(r"C:\Personal\AgentShowcase\docs\harvey\Synzo-MCP-Server-Specification.docx")


def wipe_body(doc):
    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is sectPr:
            continue
        body.remove(child)


def add_para(doc, text, style="Normal", bold=False):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if bold:
        run.bold = True
    return p


def add_bullet(doc, text):
    # Template has no List Bullet style; render as a plain paragraph with a bullet glyph.
    p = doc.add_paragraph("• " + text, style="Normal")
    return p


def add_code(doc, text):
    """Monospace, left-aligned block for JSON examples."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    return p


def add_toc(doc, heading_depth=3):
    """Insert a Word TOC field spanning Heading 1 through Heading <heading_depth>.

    The field body is empty until Word (or LibreOffice) is told to update
    fields (F9 in Word, Tools -> Update -> Update All in LibreOffice). We
    include a placeholder sentence so the reader knows to update the field
    if it looks blank.
    """
    p = doc.add_paragraph()
    run = p.add_run()
    fldChar_begin = OxmlElement("w:fldChar")
    fldChar_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fldChar_begin)

    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = f'TOC \\o "1-{heading_depth}" \\h \\z \\u'
    run._r.append(instrText)

    fldChar_sep = OxmlElement("w:fldChar")
    fldChar_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fldChar_sep)

    # Placeholder shown until the field is updated.
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose Update Field to populate the table of contents."
    run._r.append(placeholder)

    fldChar_end = OxmlElement("w:fldChar")
    fldChar_end.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar_end)
    return p


def add_kv_table(doc, pairs, header=None):
    """Two-column key/value table."""
    rows = len(pairs) + (1 if header else 0)
    table = doc.add_table(rows=rows, cols=2)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    r = 0
    if header:
        for c, text in enumerate(header):
            cell = table.rows[r].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.bold = True
        _mark_header_row_repeat(table.rows[r])
        r += 1
    for k, v in pairs:
        table.rows[r].cells[0].text = k
        table.rows[r].cells[1].text = v
        r += 1
    return table


def _mark_header_row_repeat(row):
    """Tell Word to repeat this row as a header on each new page the table spans.

    python-docx exposes this only via the underlying OOXML. Without it, long
    tables that break across pages lose their column labels on continuation.
    """
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def add_matrix_table(doc, header, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for c, text in enumerate(header):
        cell = table.rows[0].cells[c]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
    _mark_header_row_repeat(table.rows[0])
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            table.rows[r].cells[c].text = text
    return table


def add_tool_arg_table(doc, rows):
    return add_matrix_table(
        doc,
        ["Name", "Type", "Required / limits", "Description"],
        rows,
    )


def add_tool_result_table(doc, rows):
    return add_matrix_table(
        doc,
        ["Field", "Type", "Presence", "Description"],
        rows,
    )


def main():
    doc = Document(str(TEMPLATE))
    wipe_body(doc)

    # Core properties
    cp = doc.core_properties
    cp.title = "Synzo MCP Server Specification"
    cp.author = "Paul O'Hagan, Red Maple Research"
    cp.subject = "MCP connector specification for Harvey review"

    # ----- COVER -----
    add_para(doc, "Synzo MCP Server Specification", style="Title")
    add_para(
        doc,
        "Model Context Protocol connector for Synzo, a hosted document and image processing service.",
    )
    add_para(doc, "Provider: Red Maple Research")
    add_para(doc, "Primary technical contact: Paul O'Hagan, Principal — paul@redmapleresearch.ca")
    add_para(doc, "Document revision: 1.5")
    add_para(doc, "Document last updated: 5 October 2026")
    add_para(
        doc,
        "Implementation revision at this document's publication: Synzo commit 6af4183 on "
        "the master branch. The live /mcp endpoint advertises the server version "
        "(serverInfo.version, currently 0.1.0) in its initialize response, which "
        "identifies the published connector release, not the deployed Git commit. The "
        "deployed commit at any point in time is recorded in the Railway deployment log; "
        "contact paul@redmapleresearch.ca if a reviewer needs confirmation that a "
        "specific commit is in production.",
    )
    add_para(doc, "Live verification date: 5 October 2026. See Appendix B, \"Live verification record.\"")

    # ----- TABLE OF CONTENTS -----
    add_para(doc, "Contents", style="Heading 1")
    add_toc(doc, heading_depth=3)

    # ----- INTRODUCTION -----
    add_para(doc, "Introduction", style="Heading 1")
    add_para(
        doc,
        "Synzo is a hosted Model Context Protocol (MCP) connector offered by Red Maple "
        "Research. The connector exposes a five-tool processing catalog: document "
        "summarization, document translation, PII redaction, image analysis, and face "
        "obscuring. Every tool accepts the document or image bytes directly in the "
        "request as a base64-encoded argument; the server does not fetch input content "
        "from caller-supplied URLs. Billing is organization-based through metered calls.",
    )
    add_para(
        doc,
        "This document is written for engineers integrating Synzo as an MCP connector. "
        "It specifies the exact wire behavior, the authentication model, the five tools "
        "the server publishes, and the limits and controls that apply to tool execution. "
        "Where a fact depends on a provider configuration (WorkOS AuthKit) or on a "
        "deployment-time setting, that dependency is called out instead of being quoted "
        "as a hard value.",
    )
    add_para(doc, "Readers of this document should expect:")
    add_bullet(doc, "MCP protocol version and transport semantics the server actually implements.")
    add_bullet(doc, "Authentication model, including the two credential types the server accepts.")
    add_bullet(doc, "A complete catalog of the five tools, with typed arguments and typed results.")
    add_bullet(doc, "The organization-level controls that gate tool execution — Synzo does not use per-tool OAuth scopes.")
    add_bullet(doc, "Operational limits, error codes, and data-handling boundaries.")
    add_para(
        doc,
        "Many of the facts in this document are also returned dynamically through MCP "
        "discovery (initialize, tools/list, OAuth metadata). The written form here is the "
        "authoritative integration contract.",
    )

    # ----- ENDPOINT -----
    add_para(doc, "Endpoint", style="Heading 1")
    add_para(
        doc,
        "Synzo publishes one MCP endpoint over HTTPS. The transport is Streamable HTTP as "
        "defined by MCP, implemented as JSON-RPC 2.0 over HTTP POST with JSON responses. "
        "The server is stateless at the transport layer: it does not issue an "
        "Mcp-Session-Id, does not maintain a persistent server-initiated SSE stream, and "
        "does not send tool-list-change notifications.",
    )
    add_kv_table(
        doc,
        [
            ("URL", "https://www.synzo.ai/mcp"),
            ("Transport", "Streamable HTTP, JSON-RPC 2.0 over HTTP POST, JSON responses"),
            ("Server name", "synzo"),
            ("Server title", "Synzo"),
            ("Server version", "0.1.0"),
            ("Preferred MCP protocol revision", "2025-06-18"),
            ("Also accepted MCP protocol revision", "2025-03-26"),
            ("JSON-RPC version", "2.0"),
            ("tools.listChanged", "false"),
        ],
        header=["Field", "Value"],
    )
    add_para(
        doc,
        "Use Content-Type: application/json and Accept: application/json, text/event-stream "
        "on every request. GET and DELETE on /mcp return HTTP 405. The initialize method, "
        "the ping method, and tools/list are available without authentication. Any "
        "tools/call invocation requires authentication. Unknown tool names are rejected.",
    )
    add_para(
        doc,
        "The Synzo server version (0.1.0) and the MCP protocol revision are independent "
        "values. The protocol revision describes which wire specification the server "
        "complies with; the server version identifies the running implementation. Revisions "
        "to this document do not change either value.",
    )

    # ----- AUTHENTICATION -----
    add_para(doc, "Authentication", style="Heading 1")
    add_para(
        doc,
        "Synzo accepts two credential types for tool execution: OAuth 2.0 bearer tokens "
        "(Authorization Code with PKCE S256) issued through WorkOS AuthKit, and "
        "organization-scoped API keys issued through the Synzo dashboard. Both resolve to "
        "the same organization-level principal used for authorization, metering, and rate "
        "limiting. The authorization-server metadata is published under RFC 8414; Synzo "
        "does not advertise OAuth 2.1 specifically. Discovery (initialize, tools/list, "
        "OAuth metadata endpoints) is anonymous; tools/call requires authentication.",
    )
    add_para(doc, "Connection sequence:")
    add_bullet(doc, "Client POSTs initialize to https://www.synzo.ai/mcp. The server responds with its protocol version, serverInfo, and capabilities. No credential is required.")
    add_bullet(doc, "Client calls tools/list and receives the five tool descriptors. No credential is required.")
    add_bullet(doc, "Client calls tools/call. If the request is unauthenticated or the token is invalid, the server returns HTTP 401 with a WWW-Authenticate header pointing to its OAuth protected-resource metadata (JSON-RPC error -32001).")
    add_bullet(doc, "Client completes an OAuth 2.0 authorization-code flow with PKCE S256, or presents an organization-scoped API key, and retries the call.")
    add_para(doc, "OAuth discovery chain:")
    add_bullet(doc, "Protected-resource metadata: https://www.synzo.ai/.well-known/oauth-protected-resource. Its authorization_servers field lists Synzo's own public URL; the client is directed back to Synzo for the authorization-server document, not to WorkOS directly.")
    add_bullet(doc, "Authorization-server metadata: https://www.synzo.ai/.well-known/oauth-authorization-server. Synzo fetches the upstream WorkOS AuthKit metadata, pins the issuer to Synzo's configured value, injects a registration_endpoint (so clients that need DCR can find it), and sets defaults for code_challenge_methods_supported, grant_types_supported, response_types_supported, scopes_supported, and token_endpoint_auth_methods_supported where WorkOS does not already supply them. The response is an augmented view of the upstream WorkOS document, not a pure forward.")
    add_para(
        doc,
        "At the time of this document the live endpoint advertises a WorkOS AuthKit "
        "staging host (real-vine-49-staging.authkit.app) as its upstream issuer. "
        "Whether that is the intended environment for the Harvey evaluation window "
        "must be confirmed before go-live by re-reading the live authorization-server "
        "metadata.",
    )
    add_para(
        doc,
        "The augmented authorization-server metadata advertises Authorization Code with "
        "PKCE S256, refresh tokens, device-code grants, Dynamic Client Registration, and "
        "token endpoint authentication methods none, client_secret_post, and "
        "client_secret_basic. Advertising these methods in discovery does not establish "
        "that every advertised flow has been exercised end-to-end with Harvey; "
        "verification of the Harvey client registration is tracked as a pending "
        "onboarding item (see \"Open items\" in Appendix B).",
    )
    add_para(
        doc,
        "Access tokens are validated as signed JWTs against the configured issuer and "
        "audience and must carry both a sub (user) claim and an org_id (or "
        "organization_id) claim. The sub and org_id claims are used to resolve the "
        "Synzo user and organization membership; the token's email claim, if present, is "
        "used only for logging and account provisioning, not for membership resolution. "
        "API keys are accepted either in Authorization: Bearer <key> or in X-API-Key, "
        "SHA-256-hashed at rest, and compared with hmac.compare_digest in constant time. "
        "OAuth tokens must use the Bearer scheme.",
    )
    add_para(
        doc,
        "Requests containing an Origin header must match Synzo's allowlist. Harvey-"
        "specific Origin requirements will be confirmed during onboarding: whether "
        "Harvey's MCP client sends an Origin header at all depends on whether the client "
        "calls /mcp directly or proxies through a Harvey backend, and the exact value "
        "cannot be inferred from Harvey's website hostname. Requests without an Origin "
        "header remain subject to authentication and all other access controls.",
    )

    add_para(doc, "Scopes", style="Heading 2")
    add_para(
        doc,
        "Synzo does not define per-tool authorization scopes and does not filter the "
        "tool catalog by scope. All five tools use the same authentication and metering "
        "path. The identity scopes advertised by the authorization server metadata are "
        "shown below; a client should request openid plus the identity scopes it needs "
        "to establish the user's email and session.",
    )
    add_matrix_table(
        doc,
        ["Scope", "Purpose", "Notes"],
        [
            ["openid", "Required for an OpenID Connect token.", "Identity scope. No tool is gated on it; standard OIDC behavior."],
            ["profile", "Lets the authorization server return basic profile claims.", "Identity scope."],
            ["email", "Lets the authorization server return the user's email claim.", "Identity scope. Provides email for account metadata. User and organization resolution uses the sub and organization claims."],
            ["offline_access", "Lets the client receive a refresh token.", "Identity scope. Required for long-lived agent sessions."],
        ],
    )
    add_para(
        doc,
        "The scope list above is a default. If the deployed WorkOS AuthKit tenant "
        "advertises a different scopes_supported list upstream, that list is served to "
        "clients. The authoritative source is a live fetch of the authorization-server "
        "metadata URL at integration time.",
    )
    add_para(
        doc,
        "Dashboard operations (invitations, membership changes, API-key creation and "
        "revocation) are gated by a membership-role check in the Synzo application; this "
        "is independent of MCP tool execution.",
    )

    add_para(doc, "Token Security", style="Heading 2")
    add_para(
        doc,
        "Token lifetimes for OAuth access and refresh tokens are controlled by the "
        "deployed WorkOS AuthKit tenant, not by Synzo application code. The values below "
        "are read from the WorkOS AuthKit application configuration backing the submitted "
        "endpoint (Synzo application, Sessions tab) as of October 5, 2026. The password-"
        "setup and email-verification link lifetimes are not surfaced in the WorkOS "
        "dashboard; they operate at WorkOS AuthKit defaults.",
    )
    add_matrix_table(
        doc,
        ["Component", "Time-to-live (TTL)", "Notes"],
        [
            ["OAuth access token", "5 minutes", "WorkOS AuthKit \"Access token duration\" for the Synzo application. Not a Synzo-controlled value; a WorkOS configuration change here changes the lifetime without a Synzo deploy."],
            ["Session / refresh envelope", "365 days maximum, with a 2-day inactivity timeout", "WorkOS AuthKit \"Maximum session length\" (365 days) and \"Inactivity timeout\" (2 days) for the Synzo application. The refresh token issued by WorkOS is bound to this session envelope; it is renewed on use and terminates at whichever of the maximum length or inactivity timeout fires first. Revocation behavior is governed by WorkOS AuthKit; Synzo application code does not implement its own refresh-token revocation."],
            ["Organization invitation", "7 days default expiry", "WorkOS AuthKit \"Default invitation expiry\" for the Synzo application. Invitations issued through the dashboard or API expire after this interval unless an explicit override is set per invitation."],
            ["Email verification / password-setup link", "WorkOS AuthKit default", "Issued by WorkOS during account provisioning. The dashboard does not surface a configurable TTL; the links use WorkOS AuthKit's default lifetime. Treat the exact value as a WorkOS-owned implementation detail; recipients who miss the window can request a fresh link through the sign-in flow."],
            ["Synzo API key", "No time-based expiry", "A dashboard-created API key remains valid until revoked. Revocation sets a revoked_at timestamp; subsequent authentications fail with -32001. Keys are SHA-256-hashed at rest and compared with hmac.compare_digest."],
        ],
    )
    add_para(
        doc,
        "Synzo API keys are 32 bytes of secrets.token_urlsafe randomness, which produces "
        "a 256-bit key. There is no in-application lifetime check: enforcement is entirely "
        "by revocation and by the organization-level request limits.",
    )

    # ----- MCP FEATURES -----
    add_para(doc, "MCP Features", style="Heading 1")
    add_para(
        doc,
        "Synzo implements the subset of the MCP specification needed to publish and call "
        "tools. The table below enumerates which optional features are supported and which "
        "are intentionally absent.",
    )
    add_matrix_table(
        doc,
        ["Feature", "Supported", "Notes"],
        [
            ["Tools", "Yes", "Five tools are published; see Tool Catalog."],
            ["Tool input schemas", "Yes", "Published in tools/list descriptors; see also Synzo-MCP-Tool-Definitions.json."],
            ["Tool output schemas", "No", "No tool publishes an outputSchema. Return shapes are documented in the Reference section."],
            ["Structured content on results", "Yes", "Successful calls return isError: false, structuredContent with the payload, and a text content block containing the same payload serialized as indented JSON."],
            ["Resources", "No", "The server does not publish MCP resources."],
            ["Prompts", "No", "The server does not publish MCP prompts."],
            ["Roots", "No", "No roots capability is advertised."],
            ["Sampling", "No", "The server does not request client sampling."],
            ["Persistent SSE stream", "No", "Every response is a JSON HTTP response. GET on /mcp returns 405."],
            ["Mcp-Session-Id", "No", "The server does not issue session IDs."],
            ["tools.listChanged notifications", "No", "The capability is explicitly advertised as false."],
            ["Argument validation", "Handler-side only", "The dispatcher does not run general JSON Schema validation. Each handler enforces its own required-field and enum checks. Published schemas (including additionalProperties: false and the target_language 2–64 bounds) are the advertised client contract."],
        ],
    )

    # ----- TOOL CATALOG -----
    add_para(doc, "Tool Catalog", style="Heading 1")
    add_para(
        doc,
        "This section lists every tool a properly authenticated MCP client can call. All "
        "five tools are visible to every authenticated client; Synzo does not use scope-"
        "based catalog filtering. Each tool is described in detail in the Reference "
        "section.",
    )
    add_para(
        doc,
        "Every tool receives its document or image bytes directly from the client as a "
        "base64-encoded field (content_base64) in the tool arguments. The server does not "
        "fetch input content from caller-supplied URLs. A previous revision exposed an "
        "upload_file helper that returned a short-lived HTTPS URL for subsequent "
        "processing calls; that tool and the URL-input path have been removed.",
    )
    add_matrix_table(
        doc,
        ["Tool", "Purpose", "Supported input", "Result form", "Access"],
        [
            ["summarize_document", "Classify a document and produce a Markdown summary.", "Base64 bytes (content_base64) of a PDF, DOCX, PPTX, or XLSX file.", "Structured JSON with classification and summary.", "Any authenticated organization principal."],
            ["translate_document", "Translate extracted document text into a target language.", "Base64 bytes (content_base64) of a DOCX, PPTX, or XLSX file plus target_language.", "Structured JSON with translated_text (Markdown).", "Any authenticated organization principal."],
            ["redact_pii", "Produce a redacted copy of a document in its original format.", "Base64 bytes (content_base64) of a DOCX or PPTX file.", "Structured JSON with a result_url to the redacted copy.", "Any authenticated organization principal."],
            ["analyze_image", "Describe an image and flag potentially sensitive content.", "Base64 bytes (content_base64) of a JPG, JPEG, PNG, WEBP, HEIC, or HEIF image.", "Structured JSON with description, OCR text, detected objects, safety flags, and a dominant-color palette.", "Any authenticated organization principal."],
            ["detect_faces", "Blur or redact faces in an image and return a PNG.", "Base64 bytes (content_base64) of a JPG, JPEG, PNG, WEBP, HEIC, or HEIF image.", "Structured JSON with a result_url to the obscured PNG.", "Any authenticated organization principal."],
        ],
    )

    # ----- REFERENCE -----
    add_para(doc, "Reference", style="Heading 1")
    add_para(
        doc,
        "This section documents each tool in detail. Every reference page uses the same "
        "subsections: Full Description, Scopes, Call Arguments, and Return Results. A "
        "representative request and response example closes each tool. Example credentials "
        "and URLs are placeholders unless noted otherwise.",
    )
    add_para(
        doc,
        "Every tool shares one authorization rule, one success envelope, and one error "
        "classification. Rather than repeat them per tool, they are stated once here:",
    )
    add_bullet(doc, "Authorization: a valid organization-scoped principal (OAuth JWT with a provisioned organization claim, or an organization-scoped API key). No per-tool scope; no catalog filtering by scope.")
    add_bullet(doc, "Success envelope: result.isError is false; result.structuredContent is the per-tool payload described below; result.content is a one-element array of type text whose text is the same payload serialized as indented JSON.")
    add_bullet(doc, "Side effects: every successful call consumes one unit of the organization's monthly quota. redact_pii and detect_faces return a result_url to a generated binary that lives in the server's blob store for one hour. See Appendix B for retention.")
    add_para(
        doc,
        "Failures fall into two shapes. Input-validation performed inside a tool handler "
        "(unknown enum, missing or malformed base64, oversized decoded payload, provider "
        "refusal, legacy content_url argument) returns a result with isError: true and "
        "an explanatory text block. All other failures return a JSON-RPC error envelope. "
        "The mapping between HTTP status and envelope type:",
    )
    add_matrix_table(
        doc,
        ["HTTP status", "Body shape", "JSON-RPC code (if applicable)", "Notes"],
        [
            ["200", "JSON-RPC success OR error envelope", "-32002, -32003, -32004, -32005", "Successful calls, tool-level errors, organization quota/rate/unit/deadline errors."],
            ["400", "JSON-RPC error envelope", "-32700, -32600", "Invalid JSON; invalid JSON-RPC structure; request body 25-50 MB."],
            ["401", "JSON-RPC error envelope", "-32001", "Missing or invalid authentication. Carries WWW-Authenticate pointing at the protected-resource metadata URL."],
            ["404", "JSON-RPC error envelope", "-32601", "Unknown JSON-RPC method."],
            ["405", "Plain JSON error", "n/a", "Non-POST method against /mcp."],
            ["413", "JSON-RPC error envelope", "-32600", "Request body > 50 MB."],
            ["429", "Plain JSON error", "n/a", "IP-level rate limit. Honor Retry-After when present."],
        ],
    )
    add_para(
        doc,
        "Clients should read both the HTTP status and the response body on every call: "
        "a non-200 response usually also carries a JSON-RPC error envelope, and a 200 "
        "response can carry either a successful result or an isError: true tool "
        "failure.",
    )

    _write_summarize(doc, add_para, add_bullet, add_code, add_tool_arg_table, add_tool_result_table)
    _write_translate(doc, add_para, add_bullet, add_code, add_tool_arg_table, add_tool_result_table)
    _write_redact(doc, add_para, add_bullet, add_code, add_tool_arg_table, add_tool_result_table)
    _write_analyze_image(doc, add_para, add_bullet, add_code, add_tool_arg_table, add_tool_result_table)
    _write_detect_faces(doc, add_para, add_bullet, add_code, add_tool_arg_table, add_tool_result_table)

    # ----- CLIENT OPERATING GUIDANCE -----
    add_para(doc, "Client Operating Guidance", style="Heading 1")
    add_para(
        doc,
        "This section covers behavior a client needs to get right on the first "
        "integration pass: encoding the base64 argument correctly, knowing which "
        "errors should be retried and which should not, understanding what each call "
        "costs, and knowing which negative cases the server actively enforces.",
    )

    add_para(doc, "Base64 representation", style="Heading 2")
    add_para(
        doc,
        "The server validates content_base64 with Python's base64.b64decode(..., "
        "validate=True) against the standard RFC 4648 alphabet.",
    )
    add_bullet(doc, "Alphabet: standard RFC 4648 alphabet (A-Z, a-z, 0-9, + and /). The URL-safe variant (- and _) is NOT accepted.")
    add_bullet(doc, "Padding: the encoded string's length must be a multiple of 4. Standard encoders produce the right length by appending '=' padding when the input length is not a multiple of 3. Stripping that padding (as some JWT-style formats do) leaves a string whose length is not a multiple of 4; such input will be rejected. If the input length is already a multiple of 3, no padding is needed.")
    add_bullet(doc, "Line breaks: not permitted inside the string. MIME-style line-wrapped base64 (newline every 76 characters, as produced by Python's base64.encodebytes or the base64 CLI without -w0) will be rejected.")
    add_bullet(doc, "Data-URL prefixes: NOT accepted. A value beginning with 'data:application/pdf;base64,' or similar will be rejected. Strip the prefix before sending.")
    add_bullet(doc, "Whitespace: any leading, trailing, or embedded whitespace causes rejection. Trim before encoding.")
    add_bullet(doc, "Size: the decoded payload must not exceed 10,485,760 bytes. The encoded string itself is roughly 1.33x the decoded size.")
    add_para(
        doc,
        "A valid encoding of a local file is: base64.b64encode(path.read_bytes())."
        "decode('ascii') in Python, or the base64 CLI with -w0 to disable wrapping.",
    )

    add_para(doc, "First call: minimal executable example", style="Heading 2")
    add_para(
        doc,
        "The example below reads a local PDF, encodes it, calls summarize_document, "
        "and inspects both the HTTP status and the two possible failure shapes "
        "(JSON-RPC error envelope and tool-level isError). Replace SYNZO_API_KEY and "
        "the file path. The script is intended to be pasted into a Python 3.10+ "
        "environment with requests installed.",
    )
    add_code(doc,
        'import base64, json, os, sys\n'
        'import requests\n'
        '\n'
        'API_KEY = os.environ["SYNZO_API_KEY"]              # sk_synzo_...\n'
        'FILE_PATH = "sample.pdf"\n'
        'URL = "https://www.synzo.ai/mcp"\n'
        '\n'
        'with open(FILE_PATH, "rb") as f:\n'
        '    encoded = base64.b64encode(f.read()).decode("ascii")\n'
        '\n'
        'try:\n'
        '    response = requests.post(\n'
        '        URL,\n'
        '        headers={\n'
        '            "Content-Type": "application/json",\n'
        '            "Accept": "application/json, text/event-stream",\n'
        '            "MCP-Protocol-Version": "2025-06-18",\n'
        '            "Authorization": f"Bearer {API_KEY}",\n'
        '        },\n'
        '        json={\n'
        '            "jsonrpc": "2.0",\n'
        '            "id": 1,\n'
        '            "method": "tools/call",\n'
        '            "params": {\n'
        '                "name": "summarize_document",\n'
        '                "arguments": {\n'
        '                    "filename": os.path.basename(FILE_PATH),\n'
        '                    "content_base64": encoded,\n'
        '                },\n'
        '            },\n'
        '        },\n'
        '        timeout=90,\n'
        '    )\n'
        'except requests.RequestException as e:\n'
        '    # DNS/TCP/TLS failure, connection reset, read timeout, etc.\n'
        '    # Execution and charging may be unknown; see "Retry and charging."\n'
        '    print(f"Transport error: {e}", file=sys.stderr)\n'
        '    sys.exit(3)\n'
        '\n'
        'try:\n'
        '    body = response.json()\n'
        'except ValueError:\n'
        '    # A gateway (Cloudflare, Railway edge) can return HTML or plain text\n'
        '    # on 5xx without ever reaching Synzo. Treat as transport-level.\n'
        '    print(\n'
        '        f"Non-JSON response (HTTP {response.status_code}): "\n'
        '        f"{response.text[:200]!r}",\n'
        '        file=sys.stderr,\n'
        '    )\n'
        '    sys.exit(3)\n'
        '\n'
        'if "error" in body:                                 # JSON-RPC envelope error\n'
        '    print(f"JSON-RPC error (HTTP {response.status_code}):", body["error"], file=sys.stderr)\n'
        '    sys.exit(1)\n'
        '\n'
        'result = body["result"]\n'
        'if result.get("isError"):                           # tool-level error\n'
        '    message = result["content"][0]["text"]\n'
        '    print("Tool error:", message, file=sys.stderr)\n'
        '    sys.exit(2)\n'
        '\n'
        'summary = result["structuredContent"]["summary"]\n'
        'print(summary)'
    )

    add_para(doc, "Retry and charging", style="Heading 2")
    add_para(
        doc,
        "Each successful processing call consumes one call from the organization's "
        "monthly quota. Handler failures and timeouts are refunded by the metering "
        "pipeline and recorded as 'refunded' in the usage ledger. The matrix below "
        "describes which failure modes a client should retry.",
    )
    add_matrix_table(
        doc,
        ["Failure", "Retry?", "Charging", "Notes"],
        [
            ["Connection failure before the request was transmitted (DNS, TCP, TLS)", "Yes, with backoff", "Not charged", "The handler never ran. Safe to retry with exponential backoff."],
            ["Lost response after transmission (socket closed, read timeout) or HTTP 5xx from an intermediary gateway", "Yes, with backoff; see note", "May be charged", "Once the request has been transmitted, execution and charging state are unknown to the client. The handler may have run to completion; a gateway may have returned 5xx while the backend is still processing. A retry is a new request and may repeat processing and consume another unit. Confirm via the usage ledger when double-processing is unacceptable."],
            ["HTTP 429 (IP rate limit)", "Yes, after Retry-After", "Not charged", "Honor the Retry-After header."],
            ["JSON-RPC -32003 (organization rate limit, inside HTTP 200)", "Yes, after a short wait", "Not charged", "Shared across all users and tools in the organization. Back off and retry."],
            ["JSON-RPC -32005 (tool execution deadline, 60 seconds)", "Yes, but see note", "Refunded", "The deadline returns a response but does NOT forcibly terminate the worker thread. A retry can succeed even while the original worker is still running against the external provider; the retry is a fresh call and will consume a quota unit if it succeeds."],
            ["isError: true with 'blocked' or 'safety' in the message", "No", "Refunded", "Gemini safety filter refused the content. Retry with the same input will fail the same way."],
            ["isError: true with 'base64', 'content_url', or 'exceed' in the message", "No", "Refunded", "Client-side input problem. Fix the argument before retrying."],
            ["JSON-RPC -32002 (organization monthly quota exhausted)", "No, this period", "Not charged", "Resets at the start of the next calendar month, or contact support."],
            ["JSON-RPC -32001 (authentication)", "No", "Not charged", "Rotate or renew the credential."],
        ],
    )
    add_para(
        doc,
        "The server does not deduplicate requests by JSON-RPC id. Two consequences follow:",
    )
    add_bullet(doc, "After a lost response or an HTTP 5xx from an intermediary, execution and charging may be unknown. A retry is a new request that may repeat processing and consume another unit; use bounded backoff and accept the possibility of a double call when repeat processing is acceptable, or read the usage ledger before retrying when it is not.")
    add_bullet(doc, "Retries do not produce identical outputs. summarize_document, translate_document, and analyze_image are backed by Google Gemini and the model may return different wording on each call. redact_pii and detect_faces produce a new result_url and a new expires_at on each successful call. A client that depends on a specific output of an earlier call must cache that output; retrying will not reproduce it.")

    add_para(doc, "Negative verification", style="Heading 2")
    add_para(
        doc,
        "Each row below lists one rejected behavior, the expected server response, "
        "and the exact evidence it was tested against. Local evidence is an automated "
        "regression test in the repository; production evidence is a direct observation "
        "against https://www.synzo.ai/mcp. Rows that lack production evidence are not "
        "implied to be rejected in production: they are rejected by the committed "
        "handler code and verified locally against the test harness.",
    )
    add_matrix_table(
        doc,
        ["Rejected behavior", "Expected server response", "Evidence"],
        [
            ["content_url supplied alone (no content_base64)", "Tool error (isError: true) with 'content_url' in the message", "Local: tests/test_mcp_server.py::test_summarize_document_rejects_content_url_alone."],
            ["content_url and content_base64 supplied together", "Tool error (isError: true) with 'content_url' in the message; the server does NOT fall back to the base64 field", "Local: tests/test_mcp_server.py::test_summarize_document_rejects_content_url. The content_url check runs before any base64 decode."],
            ["upload_file tool present in tools/list", "Tool not advertised in the catalog", "Local: tests/test_mcp_server.py::test_upload_file_tool_is_not_advertised. Production: 2026-10-05 /mcp preflight returned exactly ['analyze_image', 'detect_faces', 'redact_pii', 'summarize_document', 'translate_document']. The cited evidence establishes catalog absence only; it does not independently establish per-tool content_url rejection on the processing tools, nor the response code a tools/call against upload_file would return (see the general error reference in Appendix A for the -32601 unknown-tool mapping)."],
            ["Base64 with invalid characters (outside the standard RFC 4648 alphabet)", "Tool error (isError: true) with 'base64' in the message", "Local: tests/test_mcp_server.py::test_processing_tool_rejects_invalid_base64 (payload '!!!not base64!!!')."],
            ["Base64 with MIME-style line breaks (newline every 76 characters)", "Tool error (isError: true) with 'base64' in the message", "Local: tests/test_mcp_server.py::test_processing_tool_rejects_base64_with_line_breaks."],
            ["Base64 of length not a multiple of 4 (stripped padding)", "Tool error (isError: true) with 'base64' in the message", "Local: tests/test_mcp_server.py::test_processing_tool_rejects_base64_missing_padding."],
            ["Decoded payload > 10 MB", "Tool error (isError: true) with 'exceed' in the message; quota refunded", "Local: tests/test_mcp_server.py::test_tools_call_decoded_content_above_10mb_returns_isError."],
            ["Missing content_base64 field (and no content_url either)", "Tool error (isError: true) with 'content_base64' in the message", "Local: tests/test_mcp_server.py::test_processing_tool_missing_content_base64_returns_isError."],
        ],
    )

    # ----- APPENDIX A: OPERATIONAL LIMITS -----
    add_para(doc, "Appendix A: Operational Limits and Errors", style="Heading 1")
    add_para(
        doc,
        "This appendix consolidates the enforced limits and the error codes the server "
        "emits. The values are read from the server configuration at the time of this "
        "revision; where a value is a deployment setting, the source is called out.",
    )
    add_matrix_table(
        doc,
        ["Limit", "Value", "Scope"],
        [
            ["Per-IP request rate", "30 requests/minute and 200 requests/hour", "All requests to /mcp, including initialize and tools/list."],
            ["Per-organization tool call rate", "Free 10/min; Starter 60/min; Pro 300/min", "Shared across all users and all tools in the organization."],
            ["Per-organization monthly tool call quota", "Free 50; Starter 10,000; Pro 100,000", "Calendar-month plan quota. The Harvey evaluation allocation is described in the Data Handling appendix."],
            ["Decoded file size cap", "10,485,760 bytes (10 MB)", "Applies to the decoded content_base64 payload on every processing tool. Rejected before processing."],
            ["Request body in the 25–50 MB range", "Rejected with HTTP 400 + JSON-RPC -32700", "Flask's MAX_CONTENT_LENGTH is 25 MB, but parsing above that limit raises an exception that the /mcp route catches as an invalid-JSON error. A reviewer-verified 26 MiB body returns HTTP 400 with a JSON-RPC envelope and code -32700 (\"Invalid JSON\"), not a bare transport error. A 10 MB file encoded as base64 expands to roughly 13.3 MB on the wire plus a small JSON envelope and fits inside the 25 MB admission limit."],
            ["Request body above 50 MB", "Rejected with HTTP 413 + JSON-RPC -32600", "The /mcp route checks request.content_length against a 50 MB hard cap before parsing. A reviewer-verified 51 MiB body returns HTTP 413 with a JSON-RPC envelope and code -32600 (\"Request body too large\")."],
            ["Tool execution deadline", "60 seconds", "Caller-facing deadline. A timed-out worker may continue running until its external operation (for example a Gemini call) returns."],
            ["Temporary generated-file lifetime", "1 hour", "Nominal expiry for binary results produced by redact_pii and detect_faces. Cleanup is lazy (performed on next access and during subsequent writes); a process restart may remove files sooner. Input documents and images are processed transiently in memory and are not staged in the blob store. The 60-second tool execution deadline returns a response to the caller but does not interrupt the worker thread; a timed-out worker (for example, waiting on a Gemini response) may continue processing and holding the decoded input bytes until its external operation completes."],
        ],
    )
    add_para(
        doc,
        "All limits apply independently. The per-IP rate cap continues to apply to higher "
        "plans. The IP-level HTTP 429 response is an ordinary JSON error object (not a "
        "JSON-RPC envelope); clients should honor the Retry-After header when present. "
        "HTTP 405 responses to non-POST methods on /mcp are similarly ordinary JSON "
        "error objects. The organization-level rate limit instead surfaces as JSON-RPC "
        "-32003 inside an HTTP 200 response.",
    )
    add_matrix_table(
        doc,
        ["JSON-RPC code", "Meaning", "HTTP correlate"],
        [
            ["-32001", "Missing or invalid authentication for a tools/call request.", "401, with WWW-Authenticate pointing at the protected-resource metadata URL."],
            ["-32002", "Organization monthly quota exhausted.", "200 with JSON-RPC error body."],
            ["-32003", "Organization rate limit exceeded.", "200 with JSON-RPC error body. (The IP-level HTTP 429 enforced at the transport layer is separate and returns an ordinary JSON error object, not a JSON-RPC envelope.)"],
            ["-32004", "Request exceeds the plan's per-call unit limit.", "200 with JSON-RPC error body."],
            ["-32005", "Overall tool execution deadline (60 seconds) exceeded.", "200 with JSON-RPC error body."],
            ["-32600", "Invalid JSON-RPC request structure (non-dict body, missing or wrong jsonrpc field, missing method).", "400."],
            ["-32601", "Unknown JSON-RPC method; also unknown tool name on tools/call.", "404 for an unknown method. An unknown tool name on tools/call returns the error inside a 200 response."],
            ["-32602", "Request-unit estimation failed. Returned by the dispatcher with the message \"Could not size request\" when a tool's units_fn raises. Handler-level invalid-parameter cases are surfaced as tool errors (isError: true), not under this code.", "200 with JSON-RPC error body."],
            ["-32700", "Request body is not valid JSON.", "400."],
        ],
    )
    add_para(
        doc,
        "The dispatcher does not run general JSON Schema validation on tool arguments. "
        "Shape, enum, and size checks happen inside each handler and surface as tool-level "
        "errors (result.isError: true) rather than as JSON-RPC -32602 errors. One "
        "dispatcher-level case still emits -32602: if the per-tool request-unit "
        "estimation (units_fn) raises, the dispatcher short-circuits with -32602 and the "
        "message \"Could not size request\". Clients should read both the HTTP status and "
        "the response body on every call: a non-200 HTTP response also carries a JSON-RPC "
        "error body in most cases, and a 200 response can still carry either a JSON-RPC "
        "error or a successful result with isError: true.",
    )

    # ----- APPENDIX B: DATA HANDLING AND EVALUATION ACCESS -----
    add_para(doc, "Appendix B: Data Handling and Evaluation Access", style="Heading 1")
    add_para(doc, "Processing flow", style="Heading 2")
    add_para(
        doc,
        "A typical tool call fans out to a small, documented set of components. Document "
        "and image summarization/translation/analysis invoke Google Gemini on the "
        "extracted text or on the decoded image bytes. PII redaction and face obscuring "
        "run locally inside the Synzo process using Microsoft Presidio with spaCy and "
        "MTCNN/OpenCV respectively; those tools do not call a third-party AI API. The "
        "server never fetches input content from a caller-supplied URL — every tool "
        "receives its document or image bytes directly in the request body.",
    )
    add_matrix_table(
        doc,
        ["Component", "Role", "Data received"],
        [
            ["MCP client (Harvey)", "Issues JSON-RPC calls over HTTPS to /mcp.", "Tool arguments: filenames, base64-encoded document or image bytes, target language, mode, blur strength."],
            ["Synzo on Railway", "Terminates HTTPS, authenticates the principal, dispatches tools, meters usage.", "Request bodies containing the inline base64 payload; the decoded bytes live in process memory for the duration of the call."],
            ["Google Gemini", "Backs summarize_document, translate_document, and analyze_image.", "Extracted document text (summarize/translate) or image bytes (analyze_image)."],
            ["Presidio + spaCy (in-process)", "Backs redact_pii.", "Extracted document text."],
            ["MTCNN + OpenCV (in-process)", "Backs detect_faces.", "Decoded image bytes."],
            ["WorkOS AuthKit", "Identity, OAuth authorization, organization membership, password and verification flows.", "User email, authentication factors, organization membership state."],
            ["PostgreSQL (Railway managed)", "Account, membership, API-key, quota, and usage-event metadata.", "Account rows, SHA-256 API-key hashes, per-call usage records (organization, tool, units, status, timestamps)."],
            ["In-process blob store", "Generated binary results only (redact_pii, detect_faces).", "Generated PNG and DOCX/PPTX result bytes. Input documents and images are never stored here."],
        ],
    )
    add_para(doc, "Isolation and the bearer-URL exception", style="Heading 2")
    add_para(
        doc,
        "Accounts, memberships, API keys, quotas, and usage events are logically scoped by "
        "organization in a shared PostgreSQL instance. Dashboard operations check "
        "membership and role. Tool execution checks that the request principal resolves to "
        "a valid organization. Isolation is enforced at the application layer; it is not "
        "based on separate databases or physical infrastructure per customer.",
    )
    add_para(
        doc,
        "Temporary file URLs returned by redact_pii and detect_faces embed a high-entropy "
        "unguessable token. Anyone holding the URL can download the file until its one-"
        "hour expiry: there is no per-organization ownership check on the blob store, and "
        "the URL is the capability. Clients must therefore treat these URLs as access "
        "credentials and avoid logging or forwarding them. These URLs are only ever "
        "returned for generated binary outputs — input documents and images are received "
        "directly as base64 and never put into the blob store.",
    )
    add_para(doc, "Retention", style="Heading 2")
    add_para(
        doc,
        "Generated binaries live in process memory and are swept on subsequent writes or "
        "on access past expiry; a process restart may evict them sooner. The one-hour "
        "value is a nominal expiry, not a guaranteed deletion deadline. Input documents "
        "and images are processed transiently in memory and are not staged in the blob "
        "store. The server's 60-second tool execution deadline returns a response to the "
        "caller but does not forcibly terminate the worker thread: a timed-out worker "
        "(for example, one still waiting for a Gemini response) may continue processing "
        "the decoded input bytes until its external operation completes. Usage events "
        "are stored in PostgreSQL with the published 90-day retention target stated in "
        "the Synzo privacy policy; automated enforcement of that period is not "
        "demonstrated in the current implementation and remains a pending operational "
        "item.",
    )
    add_para(
        doc,
        "Document and image bytes are not stored in PostgreSQL. The usage ledger records "
        "organization, authentication method, tool, units, status, error code, and "
        "timestamp. Diagnostic logs may contain error details and, on model-output parsing "
        "failures, the raw model content returned by Gemini; this package does not claim "
        "content-free logging.",
    )
    add_para(doc, "Evaluation access", style="Heading 2")
    add_para(
        doc,
        "A dedicated Harvey Connector Evaluation organization is provisioned on the "
        "production Synzo service with five member accounts: connector_eval1@harvey.ai "
        "through connector_eval5@harvey.ai. The shared allocation is 10,000 metered tool "
        "calls per calendar month across September through November 2026. September's "
        "balance was 9,994 immediately after the September 10 verification run; current "
        "balance should be confirmed with paul@redmapleresearch.ca before reporting a "
        "live number.",
    )
    add_para(
        doc,
        "To activate an evaluation account, visit https://www.synzo.ai/auth/login, request "
        "password setup for the chosen address, complete email verification, and sign in "
        "selecting Harvey Connector Evaluation when prompted. Then point the MCP client at "
        "https://www.synzo.ai/mcp and complete OAuth authorization. For additional quota, "
        "a later evaluation window, or an organization-scoped API key for scripted "
        "verification, contact paul@redmapleresearch.ca.",
    )
    add_para(doc, "Live verification record", style="Heading 2")
    add_para(
        doc,
        "A scripted end-to-end verification against the production endpoint "
        "(https://www.synzo.ai/mcp) was executed on 2026-10-05 using a temporary "
        "organization-scoped API key issued against Harvey Connector Evaluation. All "
        "five tools returned HTTP 200 with isError: false; the generated redact_pii "
        "output was downloaded and confirmed to no longer contain the synthetic email "
        "address seeded in the input; the generated detect_faces output was downloaded "
        "and confirmed to be a valid PNG. The temporary API key was revoked at the end "
        "of the run. The evaluation organization quota recorded 9,995 of 10,000 calls "
        "remaining for the current period following the verification.",
    )
    add_matrix_table(
        doc,
        ["Tool", "HTTP status", "isError", "Latency", "Additional check"],
        [
            ["summarize_document", "200", "false", "15.22 s", "structuredContent.summary present"],
            ["translate_document", "200", "false", "6.98 s", "structuredContent.translated_text present"],
            ["redact_pii", "200", "false", "0.41 s", "result_url downloaded; synthetic email absent in output"],
            ["analyze_image", "200", "false", "8.06 s", "structuredContent.analysis present"],
            ["detect_faces", "200", "false", "3.65 s", "result_url downloaded; PNG format confirmed"],
        ],
    )
    add_para(
        doc,
        "The script that produced this record is "
        "scripts/verify_harvey_evaluation.py in the Synzo repository. The run is "
        "reproducible on request against the same evaluation organization.",
    )
    add_para(doc, "Open items", style="Heading 2")
    add_para(
        doc,
        "The items below are not implied to be complete by any other section of this "
        "document. They are grouped by what resolves them: integration tests still to "
        "run between Synzo and Harvey, product capabilities that are not implemented "
        "and will not change on their own, and provider-side or operational facts that "
        "require independent confirmation outside the Synzo code base.",
    )
    add_para(doc, "Integration checks still to perform", style="Heading 3")
    add_bullet(doc, "End-to-end Harvey OAuth client registration and per-user authorization flow. The authorization-server metadata advertises PKCE S256, refresh tokens, device code, and Dynamic Client Registration, and the API-key path has been verified end-to-end (see Live verification record). The OAuth path has not been exercised with a real Harvey client.")
    add_bullet(doc, "Harvey-side attachment transfer through the content_base64 argument. The processing tools accept base64 bytes inline, but whether Harvey's MCP client implementation can read an attachment and populate that argument programmatically has not been confirmed with Harvey.")
    add_bullet(doc, "Harvey Origin allowlist entry. Whether Harvey's MCP client sends an Origin header on requests to /mcp (and if so, the exact value for evaluation and production) is not known. A Harvey backend-to-server call may send no Origin header, in which case no allowlist entry is required; a Harvey browser-direct call needs an exact Origin value added to Synzo's allowlist.")
    add_bullet(doc, "Email verification / password-setup link lifetime. The Synzo-application access token (5 minutes), session (365 days with 2-day inactivity timeout), and invitation (7 days) values were read from the WorkOS AuthKit dashboard on 2026-10-05 and are quoted in Token Security. The password-setup link lifetime is not surfaced in the dashboard and operates at the WorkOS AuthKit default; the exact value should be confirmed with WorkOS if a hard number is required.")

    add_para(doc, "Known product limitations", style="Heading 3")
    add_bullet(doc, "No Harvey-workspace-to-Synzo-organization mapping. The code base does not implement a mapping between a specific Harvey workspace identifier and a specific Synzo organization. Any Harvey user invited into a Synzo organization authenticates as a member of that organization in the normal way. A customer-admin control to restrict a Synzo organization to a particular Harvey workspace is not implemented.")
    add_bullet(doc, "No automated enforcement of the published 90-day usage-event retention target. The Synzo privacy policy states 90-day retention; the inspected implementation does not demonstrate an automated purge. Deployed cleanup remains an open operational item.")

    add_para(doc, "Provider and operational facts awaiting confirmation", style="Heading 3")
    add_bullet(doc, "Provider processing locations. Processing regions for Google Gemini, WorkOS AuthKit, and the Railway application and database are not pinned by Synzo application code. The Synzo privacy policy identifies US hosting; a complete list of provider-side processing locations has not been independently verified.")
    add_bullet(doc, "Provider retention and model-training treatment. Google Gemini content-handling terms, WorkOS data-retention terms, and Railway database retention terms as they apply to the Synzo deployment have not been independently confirmed in this package. Reviewers should obtain the relevant provider agreements rather than infer behavior from the Synzo code.")
    add_bullet(doc, "Third-party assurance. This specification does not assert SOC 2 Type II, ISO 27001, a completed third-party penetration-test report, or a dedicated prompt-injection review. A May 2026 static-analysis, dependency, and container scan report is available on request and is not equivalent to any of the above. A published Data Processing Addendum and a dedicated Trust Center are not currently identified in the submission materials.")

    # ----- REVISION HISTORY -----
    add_para(doc, "Revision History", style="Heading 1")
    add_matrix_table(
        doc,
        ["Revision", "Date", "Author", "Notes"],
        [
            ["1.0", "5 October 2026", "Paul O'Hagan, Red Maple Research", "Initial Synzo MCP Server Specification. Prepared against repository commit 1adc36e and the September 10, 2026 verification record. Supersedes the September 10 Synzo Harvey technical documentation as the integration contract."],
            ["1.1", "5 October 2026", "Paul O'Hagan, Red Maple Research", "Direct-upload revision in response to Harvey's security review. Removed the upload_file tool and the URL-input path on every processing tool; each of the five remaining tools now accepts the document or image bytes directly as content_base64. Updated Tool Catalog, every Reference entry, Operational Limits (removed URL fetch deadline and redirect cap, clarified the decoded-file scope), and Data Handling (processing-flow table, isolation statement, retention)."],
            ["1.2", "5 October 2026", "Paul O'Hagan, Red Maple Research", "Added the Live verification record subsection in Appendix B. All five tools were exercised end-to-end against the production endpoint on 2026-10-05 using a temporary API key issued against the Harvey Connector Evaluation organization; the temporary key was revoked at the end of the run. Removed the matching Outstanding verification bullet now that the live record is in place."],
            ["1.3", "5 October 2026", "Paul O'Hagan, Red Maple Research", "Review-feedback pass. Added a Client Operating Guidance section covering a runnable first-call Python example, explicit base64 encoding rules, a retry-and-charging matrix, and a negative-verification table citing the local tests and the 2026-10-05 production preflight. Added a table of contents and marked table header rows as repeating for multi-page tables. Reorganized the former 'Outstanding verification' section into integration checks still to perform, known product limitations, and provider/operational facts awaiting confirmation. Corrected cover provenance (separated implementation revision, document revision, and verification date; dropped a stale commit hash). Corrected summarize_document's filename result field to reflect Werkzeug secure_filename sanitization. Replaced the long shared error-handling paragraph with a status/body/code matrix and a one-sentence usage note. Removed per-tool repetitions of 'the server does not fetch from URLs' now that the shared bullets and the Tool Catalog intro cover it. Relabeled the redact_pii 'representative failure' as a processing limitation (the call succeeds)."],
            ["1.4", "5 October 2026", "Paul O'Hagan, Red Maple Research", "Factual corrections in the Client Operating Guidance section. Rewrote the retry matrix to distinguish failures before transmission (safe to retry; not charged) from lost responses or intermediary 5xx after transmission (execution and charging may be unknown). Replaced the 'outputs are naturally idempotent' note: Gemini-backed tools may return different wording on each call, and redact_pii/detect_faces generate new result_urls and expiries on each successful call. Narrowed every row of the negative-verification table to what the cited test actually asserts and added three new regression tests (base64 with line breaks, base64 with length not a multiple of 4, content_url supplied alone) so the table is backed by the specific checks it describes. Added transport-error handling to the runnable example (requests.RequestException and non-JSON responses). Corrected the cover provenance to clarify that /mcp advertises the server version, not the deployed Git commit, and directed reviewers to the Railway deployment log via the contact email for a specific commit confirmation. Corrected the base64 padding bullet to describe the actual rule (encoded length must be a multiple of 4) rather than imply unconditional '=' padding."],
            ["1.5", "5 October 2026", "Paul O'Hagan, Red Maple Research", "Evidence-accuracy pass on the negative-verification table. Removed the '-32601 (unknown tool)' invocation claim from the upload_file row: the cited test checks catalog absence only; the -32601 mapping is in Appendix A's general error reference and does not need to be repeated as a verified claim. Removed 'quota refunded' from the three base64-malformation rows (invalid characters, line breaks, missing padding): the cited tests assert isError and the error message only. The oversized-payload row retains 'quota refunded' because that test does explicitly verify the refund."],
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(f"Wrote: {OUTPUT}")


# ---------------------- Per-tool reference writers ----------------------

def _write_summarize(doc, add_para, add_bullet, add_code, add_arg_table, add_result_table):
    add_para(doc, "summarize_document", style="Heading 1")
    add_para(doc, "Classify a document and produce a Markdown summary.")
    add_para(doc, "Full Description", style="Heading 2")
    add_para(
        doc,
        "Extracts text from the decoded document bytes and sends the text to Google "
        "Gemini. Gemini returns an inferred document classification and a Markdown "
        "summary. Supported formats: PDF, DOCX, PPTX, XLSX. No downloadable binary is "
        "generated.",
    )
    add_para(
        doc,
        "Text extraction is best-effort for the four supported formats. Scanned PDFs may "
        "produce empty extracted text; the tool does not perform OCR on arbitrary scanned "
        "PDFs. Gemini safety filters may cause the model to decline summarization on some "
        "content; those cases are returned as a tool-level error with isError: true and an "
        "explanatory text block, not as a JSON-RPC error.",
    )
    add_para(doc, "Scopes", style="Heading 2")
    add_para(doc, "Any valid organization-scoped principal. No additional scope is required.")
    add_para(doc, "Call Arguments", style="Heading 2")
    add_arg_table(doc, [
        ["filename", "string", "Required.", "PDF, DOCX, PPTX, or XLSX filename. Used to confirm the supported extension and signature."],
        ["content_base64", "string", "Required. Decoded bytes ≤ 10,485,760 (10 MB).", "Base64-encoded document bytes. See 'Base64 representation' below for the exact rules."],
    ])
    add_para(doc, "Return Results", style="Heading 2")
    add_result_table(doc, [
        ["classification", "string | null", "May be null if the model does not return one.", "Inferred document type (e.g., Contract, Research Paper)."],
        ["summary", "string", "Always present on success.", "Markdown summary of the document."],
        ["filename", "string", "Always present.", "Input filename after Werkzeug secure_filename sanitization. Spaces, path separators, and non-ASCII characters are stripped or replaced; the extension is preserved. Not guaranteed to equal the input byte-for-byte."],
    ])
    add_para(doc, "Example request:")
    add_code(doc,
        '{\n'
        '  "jsonrpc": "2.0",\n'
        '  "id": 1,\n'
        '  "method": "tools/call",\n'
        '  "params": {\n'
        '    "name": "summarize_document",\n'
        '    "arguments": {\n'
        '      "filename": "summarize-sample.pdf",\n'
        '      "content_base64": "<base64 of the document bytes>"\n'
        '    }\n'
        '  }\n'
        '}'
    )
    add_para(doc, "Example response:")
    add_code(doc,
        '{\n'
        '  "jsonrpc": "2.0",\n'
        '  "id": 1,\n'
        '  "result": {\n'
        '    "isError": false,\n'
        '    "structuredContent": {\n'
        '      "classification": "Research Paper",\n'
        '      "summary": "## Overview\\n...",\n'
        '      "filename": "summarize-sample.pdf"\n'
        '    },\n'
        '    "content": [{"type": "text", "text": "<same payload serialized as JSON>"}]\n'
        '  }\n'
        '}'
    )
    add_para(
        doc,
        "Representative failure: a malformed content_base64 value returns a tool error "
        "(isError: true) explaining the decode failure. Supplying the legacy content_url "
        "argument returns a tool error naming the field and directing the caller to use "
        "content_base64. A decoded payload larger than 10 MB is likewise rejected as a "
        "tool error. The JSON-RPC -32005 code is reserved for the overall tool execution "
        "deadline (60 seconds) only.",
    )


def _write_translate(doc, add_para, add_bullet, add_code, add_arg_table, add_result_table):
    add_para(doc, "translate_document", style="Heading 1")
    add_para(doc, "Translate extracted document text into a target language.")
    add_para(doc, "Full Description", style="Heading 2")
    add_para(
        doc,
        "Extracts text from the decoded document bytes and sends it to Google Gemini "
        "with the requested target_language. Returns the translation as Markdown text; "
        "does not produce a translated Office file. Supported formats: DOCX, PPTX, XLSX. "
        "PDF is not supported.",
    )
    add_para(
        doc,
        "Target language is accepted as a human-readable English language name such as "
        "French, Spanish, or Simplified Chinese. The input schema advertises length bounds "
        "of 2 to 64 characters on target_language; the handler enforces presence of the "
        "field. A Gemini safety-filter refusal returns a tool error (isError: true) rather "
        "than a JSON-RPC error.",
    )
    add_para(doc, "Scopes", style="Heading 2")
    add_para(doc, "Any valid organization-scoped principal. No additional scope is required.")
    add_para(doc, "Call Arguments", style="Heading 2")
    add_arg_table(doc, [
        ["filename", "string", "Required.", "DOCX, PPTX, or XLSX filename. Used to confirm the supported extension and signature."],
        ["content_base64", "string", "Required. Decoded bytes ≤ 10,485,760 (10 MB).", "Base64-encoded document bytes. See 'Base64 representation' below for the exact rules."],
        ["target_language", "string", "Required. Advertised length 2–64.", "English language name for the destination language."],
    ])
    add_para(doc, "Return Results", style="Heading 2")
    add_result_table(doc, [
        ["filename", "string", "Always present.", "Echoed from the input."],
        ["target_language", "string", "Always present.", "Echoed from the input."],
        ["translated_text", "string", "Always present on success.", "Markdown translation of the extracted document text."],
    ])
    add_para(doc, "Example request:")
    add_code(doc,
        '{\n'
        '  "jsonrpc": "2.0",\n'
        '  "id": 2,\n'
        '  "method": "tools/call",\n'
        '  "params": {\n'
        '    "name": "translate_document",\n'
        '    "arguments": {\n'
        '      "filename": "translate-sample.docx",\n'
        '      "content_base64": "<base64 of the document bytes>",\n'
        '      "target_language": "French"\n'
        '    }\n'
        '  }\n'
        '}'
    )
    add_para(doc, "Example response:")
    add_code(doc,
        '{\n'
        '  "jsonrpc": "2.0",\n'
        '  "id": 2,\n'
        '  "result": {\n'
        '    "isError": false,\n'
        '    "structuredContent": {\n'
        '      "filename": "translate-sample.docx",\n'
        '      "target_language": "French",\n'
        '      "translated_text": "## Résumé\\n..."\n'
        '    },\n'
        '    "content": [{"type": "text", "text": "<same payload serialized as JSON>"}]\n'
        '  }\n'
        '}'
    )
    add_para(
        doc,
        "Representative failure: a PDF filename returns a tool error (isError: true) "
        "indicating the format is not supported by translate_document; use "
        "summarize_document instead.",
    )


def _write_redact(doc, add_para, add_bullet, add_code, add_arg_table, add_result_table):
    add_para(doc, "redact_pii", style="Heading 1")
    add_para(doc, "Detect and redact personally identifiable information in a document.")
    add_para(doc, "Full Description", style="Heading 2")
    add_para(
        doc,
        "Runs Microsoft Presidio with spaCy inside the Synzo process to locate "
        "personally identifiable text in the decoded document. Detected characters are "
        "replaced with block symbols. Returns a short-lived HTTPS URL (result_url) to a "
        "redacted copy of the document in its original format (DOCX or PPTX). The source "
        "file is not modified. redact_pii does not send document content to Google "
        "Gemini or to any Microsoft service.",
    )
    add_para(
        doc,
        "Detection is probabilistic and depends on the recognizers configured in Presidio "
        "and on the structural features of the input document. Text embedded in images, "
        "and text in document structures not covered by the extractor, may not be "
        "detected. Downstream review is required before disclosure; the tool does not "
        "claim complete anonymization. PDF is not supported.",
    )
    add_para(doc, "Scopes", style="Heading 2")
    add_para(doc, "Any valid organization-scoped principal. No additional scope is required.")
    add_para(doc, "Call Arguments", style="Heading 2")
    add_arg_table(doc, [
        ["filename", "string", "Required.", "DOCX or PPTX filename. Used to confirm the supported extension and signature."],
        ["content_base64", "string", "Required. Decoded bytes ≤ 10,485,760 (10 MB).", "Base64-encoded document bytes. See 'Base64 representation' below for the exact rules."],
    ])
    add_para(doc, "Return Results", style="Heading 2")
    add_result_table(doc, [
        ["filename", "string", "Always present on success.", "Name of the generated redacted copy, formed as redacted_<input filename>. Not the input filename itself."],
        ["result_url", "string (HTTPS URL)", "Always present on success.", "Bearer-access URL to the redacted copy, valid for one hour."],
        ["expires_at", "string (RFC 3339 UTC)", "Always present on success.", "Nominal expiry timestamp."],
        ["mimetype", "string", "Always present on success.", "MIME type of the redacted copy (matches the input format)."],
        ["original_size_bytes", "integer", "Always present on success.", "Size of the decoded source document."],
        ["redacted_size_bytes", "integer", "Always present on success.", "Size of the generated redacted copy."],
    ])
    add_para(doc, "Example request:")
    add_code(doc,
        '{\n'
        '  "jsonrpc": "2.0",\n'
        '  "id": 3,\n'
        '  "method": "tools/call",\n'
        '  "params": {\n'
        '    "name": "redact_pii",\n'
        '    "arguments": {\n'
        '      "filename": "redact-sample.docx",\n'
        '      "content_base64": "<base64 of the document bytes>"\n'
        '    }\n'
        '  }\n'
        '}'
    )
    add_para(doc, "Example response:")
    add_code(doc,
        '{\n'
        '  "jsonrpc": "2.0",\n'
        '  "id": 3,\n'
        '  "result": {\n'
        '    "isError": false,\n'
        '    "structuredContent": {\n'
        '      "filename": "redacted_redact-sample.docx",\n'
        '      "result_url": "https://www.synzo.ai/u/<token>",\n'
        '      "expires_at": "2026-10-05T18:30:00Z",\n'
        '      "mimetype": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",\n'
        '      "original_size_bytes": 92310,\n'
        '      "redacted_size_bytes": 93102\n'
        '    },\n'
        '    "content": [{"type": "text", "text": "<same payload serialized as JSON>"}]\n'
        '  }\n'
        '}'
    )
    add_para(
        doc,
        "Processing limitation (not a failure): a PPTX whose content is entirely inside "
        "embedded images returns a successful call whose redacted copy is visually "
        "unchanged. Text-based detection cannot see pixels. Human review remains "
        "required before disclosure.",
    )


def _write_analyze_image(doc, add_para, add_bullet, add_code, add_arg_table, add_result_table):
    add_para(doc, "analyze_image", style="Heading 1")
    add_para(doc, "Describe an image, extract visible text, and flag potentially sensitive content.")
    add_para(doc, "Full Description", style="Heading 2")
    add_para(
        doc,
        "Sends the decoded image bytes to Google Gemini vision and returns a structured "
        "JSON object with a short description, a longer rich description, OCR-extracted "
        "text, detected objects, and three safety flags. Synzo computes a dominant-color "
        "palette locally in parallel. No downloadable binary is generated.",
    )
    add_para(
        doc,
        "The analysis object is parsed from the Gemini response. Synzo does not currently "
        "publish an outputSchema for this tool, and the server does not validate the full "
        "nested analysis structure. The fields below describe the expected shape; "
        "individual values may be absent or empty when the model does not return them. The "
        "tool must be treated as model-generated: the OCR text, object list, and safety "
        "flags are useful signals, not ground truth.",
    )
    add_para(doc, "Supported formats: JPG, JPEG, PNG, WEBP, HEIC, HEIF.")
    add_para(doc, "Scopes", style="Heading 2")
    add_para(doc, "Any valid organization-scoped principal. No additional scope is required.")
    add_para(doc, "Call Arguments", style="Heading 2")
    add_arg_table(doc, [
        ["filename", "string", "Required.", "JPG, JPEG, PNG, WEBP, HEIC, or HEIF filename."],
        ["content_base64", "string", "Required. Decoded bytes ≤ 10,485,760 (10 MB).", "Base64-encoded image bytes. See 'Base64 representation' below for the exact rules."],
    ])
    add_para(doc, "Return Results", style="Heading 2")
    add_result_table(doc, [
        ["filename", "string", "Always present.", "Echoed from the input."],
        ["analysis", "object", "Always present on success.", "Nested structure described below."],
        ["dominant_colors", "array of strings", "Always present on success.", "Hexadecimal color codes (for example #3F6BAE) ordered by prominence."],
        ["analysis.description", "string", "Expected; may be empty.", "One- to two-sentence summary of the scene."],
        ["analysis.rich_description", "string", "Expected; may be empty.", "Longer narrative description."],
        ["analysis.extracted_text", "string", "Expected; may be empty.", "OCR text visible in the image."],
        ["analysis.detected_objects", "array of strings", "Expected; may be empty.", "Named objects or entities detected."],
        ["analysis.safety_flags", "object", "Expected; may be absent.", "Flags described below."],
        ["analysis.safety_flags.contains_people", "boolean", "Expected when safety_flags is present.", "Whether people are visible."],
        ["analysis.safety_flags.contains_potential_pii", "boolean", "Expected when safety_flags is present.", "Whether PII (such as visible identity documents) is suspected."],
        ["analysis.safety_flags.is_graphic_or_violent", "boolean", "Expected when safety_flags is present.", "Whether the content appears graphic or violent."],
    ])
    add_para(doc, "Example request:")
    add_code(doc,
        '{\n'
        '  "jsonrpc": "2.0",\n'
        '  "id": 4,\n'
        '  "method": "tools/call",\n'
        '  "params": {\n'
        '    "name": "analyze_image",\n'
        '    "arguments": {\n'
        '      "filename": "analyze-sample.jpg",\n'
        '      "content_base64": "<base64 of the image bytes>"\n'
        '    }\n'
        '  }\n'
        '}'
    )
    add_para(doc, "Example response:")
    add_code(doc,
        '{\n'
        '  "jsonrpc": "2.0",\n'
        '  "id": 4,\n'
        '  "result": {\n'
        '    "isError": false,\n'
        '    "structuredContent": {\n'
        '      "filename": "analyze-sample.jpg",\n'
        '      "analysis": {\n'
        '        "description": "A street scene at dusk.",\n'
        '        "rich_description": "Pedestrians cross a wet city street...",\n'
        '        "extracted_text": "ONE WAY",\n'
        '        "detected_objects": ["person", "traffic sign", "car"],\n'
        '        "safety_flags": {\n'
        '          "contains_people": true,\n'
        '          "contains_potential_pii": false,\n'
        '          "is_graphic_or_violent": false\n'
        '        }\n'
        '      },\n'
        '      "dominant_colors": ["#3F6BAE", "#141821", "#C9B57A"]\n'
        '    },\n'
        '    "content": [{"type": "text", "text": "<same payload serialized as JSON>"}]\n'
        '  }\n'
        '}'
    )
    add_para(
        doc,
        "Representative failure: a Gemini safety refusal on sensitive content returns a "
        "tool error (isError: true) explaining that the model declined, rather than a "
        "partially-populated analysis object.",
    )


def _write_detect_faces(doc, add_para, add_bullet, add_code, add_arg_table, add_result_table):
    add_para(doc, "detect_faces", style="Heading 1")
    add_para(doc, "Detect faces in an image and return a PNG with the faces blurred or redacted.")
    add_para(doc, "Full Description", style="Heading 2")
    add_para(
        doc,
        "Locates faces in the decoded image using MTCNN with OpenCV inside the Synzo "
        "process and returns a PNG in which detected faces are either blurred or covered "
        "with opaque rectangles. The source image is not modified. No face recognition "
        "is performed: detected faces are not matched to named individuals and no "
        "identity attributes are returned.",
    )
    add_para(
        doc,
        "The response does not include face counts, bounding boxes, or confidence scores. "
        "If no faces are detected, the returned PNG shows the original visual content "
        "unobscured; a client should assume this outcome is possible on any input. "
        "detect_faces does not call a third-party AI API.",
    )
    add_para(doc, "Supported formats: JPG, JPEG, PNG, WEBP, HEIC, HEIF. Output is always PNG.")
    add_para(doc, "Scopes", style="Heading 2")
    add_para(doc, "Any valid organization-scoped principal. No additional scope is required.")
    add_para(doc, "Call Arguments", style="Heading 2")
    add_arg_table(doc, [
        ["filename", "string", "Required.", "JPG, JPEG, PNG, WEBP, HEIC, or HEIF filename."],
        ["content_base64", "string", "Required. Decoded bytes ≤ 10,485,760 (10 MB).", "Base64-encoded image bytes. See 'Base64 representation' below for the exact rules."],
        ["mode", "string", "Optional. Enum: blur, redact. Default blur.", "blur applies a Gaussian blur to detected faces; redact overlays opaque rectangles."],
        ["blur_strength", "integer", "Optional. Enum: 1, 2, 3. Default 2.", "1 = light blur; 2 = strong blur; 3 = opaque fill. Ignored when mode is redact."],
    ])
    add_para(doc, "Return Results", style="Heading 2")
    add_result_table(doc, [
        ["filename", "string", "Always present on success.", "Name of the generated PNG, formed as <input basename>-faces-blurred.png when mode is blur or <input basename>-faces-redacted.png when mode is redact. Not the input filename itself."],
        ["mode", "string", "Always present.", "Echoed from the input (or the default)."],
        ["result_url", "string (HTTPS URL)", "Always present on success.", "Bearer-access URL to the generated PNG, valid for one hour."],
        ["expires_at", "string (RFC 3339 UTC)", "Always present on success.", "Nominal expiry timestamp."],
        ["mimetype", "string", "Always present on success.", "Always image/png."],
    ])
    add_para(doc, "Example request:")
    add_code(doc,
        '{\n'
        '  "jsonrpc": "2.0",\n'
        '  "id": 5,\n'
        '  "method": "tools/call",\n'
        '  "params": {\n'
        '    "name": "detect_faces",\n'
        '    "arguments": {\n'
        '      "filename": "detect-faces-sample.jpg",\n'
        '      "content_base64": "<base64 of the image bytes>",\n'
        '      "mode": "redact"\n'
        '    }\n'
        '  }\n'
        '}'
    )
    add_para(doc, "Example response:")
    add_code(doc,
        '{\n'
        '  "jsonrpc": "2.0",\n'
        '  "id": 5,\n'
        '  "result": {\n'
        '    "isError": false,\n'
        '    "structuredContent": {\n'
        '      "filename": "detect-faces-sample-faces-redacted.png",\n'
        '      "mode": "redact",\n'
        '      "result_url": "https://www.synzo.ai/u/<token>",\n'
        '      "expires_at": "2026-10-05T18:45:00Z",\n'
        '      "mimetype": "image/png"\n'
        '    },\n'
        '    "content": [{"type": "text", "text": "<same payload serialized as JSON>"}]\n'
        '  }\n'
        '}'
    )
    add_para(
        doc,
        "Representative failure: an unknown mode value is rejected before processing with "
        "a tool-level error (result.isError: true) identifying the invalid enum. A "
        "blur_strength value outside {1, 2, 3} returns the same tool-level error. These "
        "are tool errors, not JSON-RPC -32602 errors: the server's dispatcher does not run "
        "general JSON Schema validation, so argument-shape checks happen inside each "
        "handler.",
    )


if __name__ == "__main__":
    main()
