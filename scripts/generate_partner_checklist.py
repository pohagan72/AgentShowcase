"""Generate an Excel checklist for partners submitting an MCP server to the Anthropic Connector Directory.

Built from lessons learned doing Synzo's real 2026-06 submission. Product-neutral —
no Synzo-specific values in the user-facing output. Writes to the user's Downloads
folder.

Run:
    .venv/Scripts/python -m scripts.generate_partner_checklist
"""

from __future__ import annotations

import os
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


# ---------- Styles ----------

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
SECTION_FILL = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
SECTION_FONT = Font(bold=True, size=12, color="1E3A8A")
LESSON_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
CRITICAL_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
TIP_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

THIN = Side(border_style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ALIGN_WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_header_row(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER


def style_section_row(ws, row: int, ncols: int, text: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.fill = SECTION_FILL
    cell.font = SECTION_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left", indent=1)
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).border = BORDER


def style_data_row(ws, row: int, ncols: int, fill=None) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.alignment = ALIGN_WRAP
        cell.border = BORDER
        if fill:
            cell.fill = fill


def add_yesno_validation(ws, cell_range: str) -> None:
    dv = DataValidation(type="list", formula1='"Yes,No,N/A"', allow_blank=True)
    dv.add(cell_range)
    ws.add_data_validation(dv)


# ---------- Tab 1: Read me first ----------


def build_tab_readme(wb) -> None:
    ws = wb.active
    ws.title = "1. Read Me First"

    ws["A1"] = "MCP Connector Directory Submission — Partner Checklist"
    ws["A1"].font = Font(bold=True, size=16, color="1E3A8A")
    ws.merge_cells("A1:B1")

    sections = [
        (
            "What this is",
            "A walkthrough of every step needed to submit an MCP server to the Anthropic Connector Directory. "
            "Built from lessons learned doing a real submission in 2026-06. Use the tabs in order — each one "
            "is a discrete phase of work.",
        ),
        (
            "Who this is for",
            "Developers and product owners who have built an MCP server and want to publish it in claude.ai's "
            "Connector Directory so end users can install it with one click.",
        ),
        (
            "Time estimate",
            "If your server is already built and deployed: ~4-6 hours total for first-time submitters. Most of "
            "that is documentation, screenshots, and form-filling — not coding. If your server isn't built yet, "
            "budget 2-4 weeks for pre-submission work.",
        ),
        (
            "Tab order",
            "1. Read Me First (this page)\n"
            "2. Pre-flight Checklist (run every check before opening the form)\n"
            "3. Public Site Requirements (docs / privacy / support / terms)\n"
            "4. Auth & Test Account (the highest-stakes section)\n"
            "5. Screenshots (3-5 promotional images for the listing)\n"
            "6. Form Page-by-Page (the actual form, field by field)\n"
            "7. Paste-Ready Templates (boilerplate text to customize)\n"
            "8. Post-Submission (what to do after clicking Submit)",
        ),
        (
            "How to use it",
            "(1) Read this page top to bottom.\n"
            "(2) Walk Tab 2 — every row must be Yes or N/A before opening the form. Use the dropdown in the Done? column.\n"
            "(3) Tabs 3-5 — build the assets you'll link to in the form. Don't open the form until these are live.\n"
            "(4) Tab 6 — fill the form, one section at a time. Use the Your Answer column to draft, then paste into Google Forms.\n"
            "(5) Tab 7 — paste-ready boilerplate text. Customize the placeholders (PRODUCT_NAME, API_KEY, etc.).\n"
            "(6) Tab 8 — what to do after submitting. Includes follow-up tracking.",
        ),
        (
            "Critical principles",
            "• The form is unforgiving — once submitted, you can't edit fields. Triple-check before clicking Submit.\n"
            "• Reviewer-instructions field is the highest-leverage field in the entire form. Treat it as more "
            "important than any single checkbox.\n"
            "• Every URL you put in the form will be tested by a reviewer. Test them yourself with curl first. "
            "Don't skip this — broken URLs are the #1 avoidable mistake.\n"
            "• Screenshots must show ACTUAL TOOL OUTPUT, not mid-execution permission prompts or setup pages.\n"
            "• The form's auth options are limited and may not match your real setup. Pick the closest match "
            "and explain the real story in the reviewer-instructions field.",
        ),
        (
            "What you need ready before starting",
            "• A live, deployed MCP server (HTTPS endpoint).\n"
            "• A test account on your service with credentials that don't expire for 30+ days.\n"
            "• A logo (SVG preferred, 1:1 square, ≥ 500×500). PNG accepted with caveats.\n"
            "• 3-5 promotional screenshots (≥ 1000 px wide, PNG).\n"
            "• Public-facing pages: docs, privacy policy, support, terms.\n"
            "• An email address you'll monitor for reviewer correspondence (typical review: 2-4 weeks).",
        ),
        (
            "Output of this exercise",
            "A completed Anthropic Connector Directory submission, in the review queue. Typical review time is "
            "2-4 weeks. Reviewer correspondence happens via the primary contact email on the form. If approved, "
            "your connector appears in claude.ai's Connector Directory.",
        ),
        (
            "Color key used throughout this workbook",
            "BLUE bar = section header (organizational)\n"
            "RED cell = critical / blocker item (don't skip)\n"
            "YELLOW cell = lesson learned (avoid this mistake)\n"
            "GREEN cell = tip / nice-to-have (improves submission quality)",
        ),
    ]

    for i, (heading, text) in enumerate(sections, start=3):
        h_cell = ws.cell(row=i, column=1, value=heading)
        h_cell.font = Font(bold=True, color="1E3A8A")
        h_cell.alignment = Alignment(wrap_text=True, vertical="top")
        h_cell.fill = SUBHEADER_FILL
        h_cell.border = BORDER
        t_cell = ws.cell(row=i, column=2, value=text)
        t_cell.alignment = ALIGN_WRAP
        t_cell.border = BORDER

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    for r in range(3, 3 + len(sections)):
        ws.row_dimensions[r].height = 95


# ---------- Tab 2: Pre-flight checklist ----------


def build_tab_preflight(wb) -> None:
    ws = wb.create_sheet("2. Pre-flight Checklist")
    headers = ["#", "Done?", "Item", "Why it matters", "How to verify"]
    ws.append(headers)
    style_header_row(ws, 1, 5)

    SECTION = object()

    rows = [
        (SECTION, "Server is healthy"),
        ("Production endpoint returns 200 to a basic POST tools/list",
         "If the reviewer's first request fails, they may abandon the review.",
         'curl -sS -X POST https://YOUR-DOMAIN/mcp -H "Content-Type: application/json" -d \'{"jsonrpc":"2.0","id":1,"method":"tools/list"}\''),
        ("All advertised tools appear in the tools/list response",
         "Drift between your docs and your live server is a fast bounce.",
         "Compare the tools/list output to your docs page."),
        ("Server runs on HTTPS (not HTTP)",
         "Required by the form. http:// gets rejected.",
         "Check the URL scheme; confirm your host terminates TLS at the edge."),
        ("CORS headers correctly allow claude.ai origin",
         "Browser-based MCP clients need this. Missing CORS = silent failures.",
         'curl -i -X OPTIONS https://YOUR-DOMAIN/mcp -H "Origin: https://claude.ai" — should return 204 with Access-Control-Allow-Origin: https://claude.ai'),
        ("Server returns 401 + WWW-Authenticate on unauthenticated tools/call",
         "Required by MCP spec §2.1 + RFC 9728 so OAuth clients can discover the auth server.",
         "curl -i -X POST https://YOUR-DOMAIN/mcp with an unauthenticated tools/call body"),

        (SECTION, "Auth is working"),
        ("Test credentials work end-to-end against production",
         "Reviewer pastes them and expects it to just work.",
         "Run a tools/call with your test credentials; confirm 200 + correct result."),
        ("Test credentials valid for 30+ days from submission",
         "Form requires this. Don't generate a key that expires in 7 days.",
         "Check expiry of the key/token/credential you're providing."),
        ("No 2FA required for test account",
         "Reviewer can't pass 2FA challenges. If you need MFA, use mcp-review@anthropic.com as test email.",
         "Open an incognito browser and sign in as the test account; confirm no 2FA prompt."),
        ("Reviewer org has enough quota for a 5-call sweep + testing slack",
         "If reviewer's sweep exhausts quota, they can't complete review.",
         "Check your dashboard's quota counter. Aim for ≥ 15 calls headroom."),
        ("Auth flow you're submitting (OAuth / API key / etc.) matches the form's options",
         "Form's auth options are limited; you may need to pick the closest match.",
         "Read Tab 4 for the full auth-matching playbook."),

        (SECTION, "Public URLs"),
        ("Documentation URL returns 200, publicly accessible (no login required)",
         "Form requires a public docs URL.",
         "curl -sI https://YOUR-DOMAIN/docs — should be 200."),
        ("Privacy policy URL returns 200",
         "Required field. Missing this blocks submission.",
         "curl -sI https://YOUR-DOMAIN/privacy"),
        ("Support URL returns 200",
         "Required field. Must be distinct from docs URL.",
         "curl -sI https://YOUR-DOMAIN/support"),
        ("Terms of service URL returns 200",
         "Required by the form's Documentation checklist.",
         "curl -sI https://YOUR-DOMAIN/terms"),
        ("CRITICAL: Logo URL returns 200 with correct Content-Type",
         "CRITICAL: a 2026-06 submission had a logo URL that returned 404 because the file was on disk but not deployed. Reviewer saw a broken logo.",
         "curl -sI https://YOUR-DOMAIN/path/to/logo.svg — must return 200 + image/svg+xml or image/png"),
        ("Sample-file URLs (if you host any) all return 200",
         "If your reviewer-instructions tell the reviewer to use these URLs, they MUST work.",
         "curl -sI each URL one by one."),
        ("All URLs were tested AFTER your most recent deploy, not before",
         "Pre-deploy testing doesn't prove a URL is live.",
         "Check deploy log timestamps; verify with curl from a clean machine."),

        (SECTION, "Documentation content"),
        ("Docs page lists every tool with description + example prompt",
         "Reviewer compares against the form's server-inventory field.",
         "Visually inspect your /docs page."),
        ("Docs include setup instructions for at least claude.ai (web)",
         "Form has a checkbox attesting to this.",
         'Look at your /docs page — is there a "How to connect" section?'),
        ("Docs include a troubleshooting section",
         "Form has a checkbox attesting to this.",
         "Visually inspect."),
        ("Privacy policy is GDPR-scoped (if accepting EU users)",
         "Form has a GDPR-compliance checkbox.",
         "Read your /privacy page — does it cover lawful basis, data subject rights, international transfers?"),
        ("Privacy policy names all third-party processors (sub-processors)",
         "Required by GDPR + good practice.",
         "List every third-party service your tools touch (AI model providers, hosting, etc.)."),

        (SECTION, "Code-side hygiene (per Anthropic Software Directory Policy)"),
        ('Every tool has a non-empty "title" field in its registration',
         'Form has a checkbox attesting to "user-friendly titles for all tools".',
         'Inspect your tools/list response — every tool should have a "title".'),
        ("Every tool has all 4 MCP annotations: readOnlyHint, destructiveHint, idempotentHint, openWorldHint",
         'Form has a checkbox attesting to "accurate tool annotations".',
         'Inspect your tools/list response — every tool should have all 4 annotations set.'),
        ("Tool descriptions don't instruct Claude to call other tools or override system instructions",
         "Anthropic Policy 2 (prompt-injection) blocker.",
         'Read each description out loud. Does it say "Claude should also call..." or "ignore system instructions"? If yes, fix it.'),
        ("No observability/analytics SDKs that send tool data off-platform",
         "Anthropic Policy 1.D / 1.F. Sentry, Datadog, PostHog, GA, etc. capturing tool inputs/outputs would block.",
         "grep your codebase for sentry/datadog/posthog/etc.; confirm none are wired up."),
        ("Error messages are actionable (not generic 'Internal Server Error')",
         "Anthropic Policy 5.A. Reviewer will trip an error case during testing.",
         "Force a quota-exceeded / bad-input case and inspect the error response — does it tell the caller what went wrong?"),
        ("Tool responses are minimal — no echoing of the source data back",
         "Anthropic Policy 5.B (token frugality). Returning the entire source document inflates costs.",
         "Inspect a tools/call response; should be just the structured result, not source-text echo."),

        (SECTION, "Test in the actual chat host"),
        ("Verified at least ONE tool call works end-to-end inside claude.ai web",
         "You need to know what reviewer will see. Surprises here block submission.",
         "Add your connector to claude.ai, run one prompt, confirm clean result. Take a screenshot for Tab 5."),
        ("Tool calls complete in under 30 seconds for typical inputs",
         "Slower tools risk timing out the chat host.",
         "Time a real call end-to-end with a stopwatch."),
        ("No 'base64 stalling' in claude.ai (chat sandbox visibly building base64 string)",
         "If users see the chat stall for minutes constructing base64, the connector feels broken. Common issue when tools take inline base64 inputs.",
         "Watch the chat-sandbox panel. If it spends a long time on 'Embedding base64...' or 'Resize to make base64 smaller', switch to URL-based inputs or add an upload tool that returns a URL."),

        (SECTION, "Submission materials"),
        ("Logo file ready (SVG preferred, 1:1, ≥ 500×500)",
         "Required by form. SVG is preferred but PNG accepted in practice.",
         "Have the file ready and hosted at a public URL."),
        ("3-5 screenshots captured at ≥ 1000 px wide",
         "Form requires this. Captures below 1000 px will be rejected.",
         "Use the dimension-check script in Tab 5 (or right-click → Properties on Windows)."),
        ("Screenshots show actual tool output, not mid-execution states",
         "A mid-execution 'Allow / Deny' permission prompt looks like the tool is broken in screenshots.",
         "Click 'Always allow' first, wait for full response, THEN capture."),
        ("Tagline ≤ 55 characters drafted",
         "Form-required field. Char-counts include spaces.",
         "Count characters; aim for 45-50 to be safe."),
        ("Description 50-100 words drafted",
         "Form-required field with strict word count.",
         "Use a word counter. Stay between 60 and 90 to be safe."),
        ("3+ use cases drafted with example prompts",
         "Form-required field. Reviewer uses these to test.",
         "Each use case = 1-line description + paste-ready prompt. URL-based prompts are best."),
        ("Pre-submission checklist in Tab 6 'Form Page-by-Page' walked top-to-bottom in draft form",
         "Catching missing fields before opening the form saves time.",
         "Use the Your Answer column on Tab 6 to draft every answer before opening Google Forms."),
    ]

    row = 2
    item_idx = 0
    for entry in rows:
        if entry[0] is SECTION:
            style_section_row(ws, row, 5, entry[1])
            row += 1
            continue
        item, why, how = entry
        item_idx += 1
        ws.cell(row=row, column=1, value=item_idx)
        ws.cell(row=row, column=2, value="").alignment = ALIGN_CENTER
        ws.cell(row=row, column=3, value=item)
        ws.cell(row=row, column=4, value=why)
        ws.cell(row=row, column=5, value=how)
        style_data_row(ws, row, 5)
        if item.startswith("CRITICAL") or "CRITICAL" in why:
            for c in range(3, 6):
                ws.cell(row=row, column=c).fill = CRITICAL_FILL
        row += 1

    widths = [5, 9, 52, 55, 65]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, row):
        ws.row_dimensions[r].height = 42

    add_yesno_validation(ws, f"B2:B{row-1}")
    ws.freeze_panes = "A2"


# ---------- Tab 3: Public site requirements ----------


def build_tab_public_site(wb) -> None:
    ws = wb.create_sheet("3. Public Site Requirements")
    headers = ["Page", "Required content", "URL to provide on form", "Your URL", "Done?"]
    ws.append(headers)
    style_header_row(ws, 1, 5)

    rows = [
        ("/docs (or equivalent)",
         "Server description, list of all tools (name + description + example prompt), setup instructions for "
         "claude.ai (web) and Claude Desktop, troubleshooting guide referencing your error codes, and a FAQ.",
         "https://YOUR-DOMAIN/docs",
         "", ""),
        ("/privacy",
         "GDPR-compliant if accepting EU users. Must cover: data collection (account, usage, log, cookies), "
         "data use, third-party processors (every service your tools call — name them explicitly), data "
         "retention windows, international transfers (SCCs if applicable), security, data subject rights, "
         "minors policy, contact for data requests.",
         "https://YOUR-DOMAIN/privacy",
         "", ""),
        ("/support",
         "Contact mailbox, response SLA, what info to include in a support report, link to /docs for "
         "self-serve, security disclosure path.",
         "https://YOUR-DOMAIN/support",
         "", ""),
        ("/terms (or /tos)",
         "Standard terms of service: acceptable use, account responsibilities, intellectual property, "
         "warranty disclaimer, limitation of liability, governing law.",
         "https://YOUR-DOMAIN/terms",
         "", ""),
        ("/security (optional but recommended)",
         "Security vulnerability disclosure process. Required by Anthropic Software Directory Terms to have "
         "'a mechanism for receiving reports of security vulnerabilities'.",
         "https://YOUR-DOMAIN/security",
         "", ""),
        ("Logo at a stable URL",
         "Square (1:1), SVG preferred, ≥ 500×500. PNG accepted in practice. Hosted at a URL that won't change.",
         "https://YOUR-DOMAIN/path/to/logo.svg",
         "", ""),
        ("Sample files for reviewer (highly recommended)",
         "If your tools take file inputs, host 3-5 sample files at public URLs. This lets reviewer test without "
         "uploading anything. Use these in your use-case prompts. See Tab 4 'Sample files' section for guidance.",
         "https://YOUR-DOMAIN/static/samples/* (or similar)",
         "", ""),
    ]

    row = 2
    for page, content, url, _, _ in rows:
        ws.cell(row=row, column=1, value=page)
        ws.cell(row=row, column=2, value=content)
        ws.cell(row=row, column=3, value=url)
        ws.cell(row=row, column=4, value="")
        ws.cell(row=row, column=5, value="")
        style_data_row(ws, row, 5)
        row += 1

    widths = [30, 60, 45, 40, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, row):
        ws.row_dimensions[r].height = 90

    add_yesno_validation(ws, f"E2:E{row-1}")
    ws.freeze_panes = "A2"

    # Lessons callout
    lesson_row = row + 1
    ws.cell(row=lesson_row, column=1, value="LESSONS LEARNED").font = Font(bold=True, color="92400E")
    ws.cell(row=lesson_row, column=1).fill = LESSON_FILL
    ws.merge_cells(start_row=lesson_row, start_column=1, end_row=lesson_row, end_column=5)
    lesson_row += 1
    lessons = [
        "Privacy policy: name EVERY downstream service that touches user data. Reviewers check this against your "
        "code claims. If you say 'no third-party trackers' make sure you actually have none.",
        "Sample files: hosting reviewer samples at public URLs is the single biggest UX improvement you can make "
        "to the reviewer experience. With URL-input tools, the chat sandbox stays fast.",
        "Logo: SVG is preferred but auto-traced SVGs (from PNG) are acceptable. ~30 minutes with an online "
        "vectorizer. Quality is good enough for directory thumbnails.",
        "URLs that don't return 200 at submission time = broken listing. Test EVERY URL with curl right before submitting.",
    ]
    for lesson in lessons:
        ws.cell(row=lesson_row, column=1, value="• " + lesson)
        ws.cell(row=lesson_row, column=1).fill = LESSON_FILL
        ws.cell(row=lesson_row, column=1).alignment = ALIGN_WRAP
        ws.merge_cells(start_row=lesson_row, start_column=1, end_row=lesson_row, end_column=5)
        ws.row_dimensions[lesson_row].height = 45
        lesson_row += 1


# ---------- Tab 4: Auth & test account ----------


def build_tab_auth(wb) -> None:
    ws = wb.create_sheet("4. Auth & Test Account")

    # Top: critical context
    ws["A1"] = "Authentication & Test Account Setup"
    ws["A1"].font = Font(bold=True, size=14, color="1E3A8A")
    ws.merge_cells("A1:D1")

    ws["A2"] = (
        "This is the highest-stakes section of the submission. Reviewer will test your auth flow and your "
        "tool sweep. If credentials don't work, or quota is exhausted, or the auth setup is confusing, "
        "your submission gets bounced."
    )
    ws["A2"].alignment = ALIGN_WRAP
    ws["A2"].fill = CRITICAL_FILL
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 50

    # Auth-options matrix
    ws["A4"] = "Matching your auth flow to the form's options"
    ws["A4"].font = SECTION_FONT
    ws["A4"].fill = SECTION_FILL
    ws.merge_cells("A4:D4")
    ws.row_dimensions[4].height = 22

    auth_table_headers = ["Your actual auth", "Pick on form: Authentication Type", "Pick on form: Auth Client", "Notes"]
    for i, h in enumerate(auth_table_headers, 1):
        ws.cell(row=5, column=i, value=h)
    style_header_row(ws, 5, 4)

    auth_rows = [
        ("OAuth 2.0 with Dynamic Client Registration (DCR)",
         "OAuth 2.0",
         "Dynamic OAuth Client (DCR / CIMD)",
         "Cleanest fit. Form's logic matches your reality. Static Client ID / Secret fields stay blank."),
        ("OAuth 2.0 with a pre-registered static client",
         "OAuth 2.0",
         "Static OAuth Client",
         "Fill Static Client ID + Secret with the pre-registered credentials."),
        ("API key (Bearer token) — NOT OAuth",
         "OAuth 2.0",
         "Dynamic OAuth Client (DCR / CIMD)",
         "MISMATCH: form has no 'API key' option. Pick OAuth 2.0 (closest) and explain in reviewer-instructions. "
         "Leave Static Client ID / Secret blank. The reviewer-instructions field is where the API key actually goes."),
        ("No authentication required",
         "No auth needed",
         "(N/A)",
         "Rare for MCP servers but supported. Leave reviewer-instructions noting that no credentials are needed."),
    ]
    row = 6
    for r in auth_rows:
        for c, val in enumerate(r, 1):
            ws.cell(row=row, column=c, value=val)
        style_data_row(ws, row, 4)
        if "MISMATCH" in r[3]:
            for c in range(1, 5):
                ws.cell(row=row, column=c).fill = CRITICAL_FILL
        ws.row_dimensions[row].height = 70
        row += 1

    # Test account requirements
    row += 1
    ws.cell(row=row, column=1, value="Test account requirements").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 22
    row += 1

    req_headers = ["Requirement", "Why", "How to satisfy", "Done?"]
    for i, h in enumerate(req_headers, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, 4)
    row += 1

    requirements = [
        ("Credentials valid 30+ days from submission",
         "Form-required attestation. Reviewer may not test for 2-4 weeks.",
         "Generate a long-lived API key / set auth-token TTL to 30+ days. Don't revoke during review."),
        ("Free of 2FA / MFA challenges",
         "Reviewer can't pass MFA. Hard blocker.",
         "Use a dedicated test account with MFA disabled, OR use mcp-review@anthropic.com if MFA is mandatory."),
        ("Bound to a sample data / quota set ready for testing",
         "Reviewer will run your tool sweep. Empty account or zero quota blocks the test.",
         "Pre-seed the test org with sample data. Confirm quota counter ≥ 15 before submitting."),
        ("Credentials documented clearly in the reviewer-instructions field",
         "Reviewer needs to know HOW to use the credentials (which header, which field).",
         "See Tab 7 'Paste-Ready Templates' for the reviewer-instructions boilerplate."),
        ("Credentials are placeholdered, not pasted, in this workbook",
         "This workbook may be shared with collaborators. Don't paste real secrets here.",
         "Use <PASTE_AT_SUBMISSION_TIME> placeholders; paste real values into Google Forms only."),
    ]
    for req, why, how in requirements:
        ws.cell(row=row, column=1, value=req)
        ws.cell(row=row, column=2, value=why)
        ws.cell(row=row, column=3, value=how)
        ws.cell(row=row, column=4, value="")
        style_data_row(ws, row, 4)
        ws.row_dimensions[row].height = 55
        row += 1
    add_yesno_validation(ws, f"D{row-len(requirements)}:D{row-1}")

    # Sample files section
    row += 1
    ws.cell(row=row, column=1, value="Sample files (highly recommended)").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 22
    row += 1

    sample_intro = (
        "If your tools take file inputs (PDF / DOCX / image / etc.), host 3-5 sample files at public URLs on "
        "your domain. Reviewers can then test by referencing the URLs in chat prompts — no file upload needed. "
        "This is the single biggest UX win for the reviewer experience. Also lets you write use-case prompts "
        "that work without context: \"Use YourProduct to summarize https://your-domain/static/samples/sample.pdf\"."
    )
    ws.cell(row=row, column=1, value=sample_intro)
    ws.cell(row=row, column=1).alignment = ALIGN_WRAP
    ws.cell(row=row, column=1).fill = TIP_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 75
    row += 2

    sample_headers = ["Tool type", "Sample file ideas", "Source (use public-domain or synthetic)", "Sample URL planned"]
    for i, h in enumerate(sample_headers, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, 4)
    row += 1

    samples = [
        ("Text doc tool (summarize / translate)",
         "Short story or article (public domain), corporate report (synthetic), academic paper preprint.",
         "Project Gutenberg, US government PDFs, your own synthesized doc. AVOID in-copyright work.",
         ""),
        ("PII / data extraction tool",
         "Synthetic HR memo or contract with seeded fake PII (fake names, fake SSNs, fake emails).",
         "Generate yourself with python-docx. NEVER use real PII.",
         ""),
        ("Image analysis tool",
         "Scenic photo with visible text, multi-object scene, or a stock photo.",
         "Your own photos, Unsplash, public domain.",
         ""),
        ("Face detection / image processing tool",
         "Group photo (small number of clearly visible faces).",
         "Permissioned photo or stock photo with model release.",
         ""),
        ("Other modality (audio / video)",
         "Short clip (≤ 30s typically).",
         "Public-domain or your own creation.",
         ""),
    ]
    for s in samples:
        for c, val in enumerate(s, 1):
            ws.cell(row=row, column=c, value=val)
        style_data_row(ws, row, 4)
        ws.row_dimensions[row].height = 55
        row += 1

    widths = [38, 38, 45, 38]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------- Tab 5: Screenshots ----------


def build_tab_screenshots(wb) -> None:
    ws = wb.create_sheet("5. Screenshots")

    ws["A1"] = "Promotional Screenshots"
    ws["A1"].font = Font(bold=True, size=14, color="1E3A8A")
    ws.merge_cells("A1:E1")

    ws["A2"] = (
        "Form requires 3-5 promotional screenshots showing your connector running inside claude.ai. "
        "Each must be ≥ 1000 px wide, PNG format, cropped to the response. Each should pair with the "
        "prompt that produced it."
    )
    ws["A2"].alignment = ALIGN_WRAP
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 50

    # Hard requirements
    ws["A4"] = "Hard requirements (form will reject if missing)"
    ws["A4"].font = SECTION_FONT
    ws["A4"].fill = SECTION_FILL
    ws.merge_cells("A4:E4")
    ws.row_dimensions[4].height = 22

    reqs = [
        ("Width", "≥ 1000 pixels", 'Right-click file → Properties → Details, or run: python -c "from PIL import Image; print(Image.open(\'path\').size)"'),
        ("Format", "PNG", "JPG / WEBP not accepted by the form."),
        ("Aspect", "No constraint — landscape is fine, portrait is fine.", "Whatever cleanly shows the prompt + response."),
        ("Content", "Tool's ACTUAL output, not mid-execution states", "Don't capture at the 'Always allow / Deny' permission prompt. Click Allow first, wait for response, THEN capture."),
        ("Crop", "Tight to the response, no full window chrome", "Crop out the browser address bar, the claude.ai sidebar, the message-input box at the bottom."),
        ("Prompt visibility", "The prompt that produced the response should be visible in the screenshot or paired in the form's caption", "Capture from just above the prompt down to just below the structured result."),
    ]
    row = 5
    for i, h in enumerate(["What", "Spec", "How to verify"], 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, 3)
    row += 1
    for r in reqs:
        for c, val in enumerate(r, 1):
            ws.cell(row=row, column=c, value=val)
        style_data_row(ws, row, 3)
        ws.row_dimensions[row].height = 50
        row += 1

    # Per-screenshot tracker
    row += 1
    ws.cell(row=row, column=1, value="Per-screenshot tracker").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.row_dimensions[row].height = 22
    row += 1

    headers = ["#", "Tool / use case", "Prompt that produced it", "File path / dimensions", "Done?"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, 5)
    row += 1
    start_track = row
    for n in range(1, 6):
        ws.cell(row=row, column=1, value=n)
        ws.cell(row=row, column=2, value="<TOOL_NAME or USE_CASE>")
        ws.cell(row=row, column=3, value="<PASTE_PROMPT_HERE>")
        ws.cell(row=row, column=4, value="<file path + WxH dimensions>")
        ws.cell(row=row, column=5, value="")
        style_data_row(ws, row, 5)
        ws.row_dimensions[row].height = 50
        row += 1
    add_yesno_validation(ws, f"E{start_track}:E{row-1}")

    # Capture workflow
    row += 1
    ws.cell(row=row, column=1, value="Capture workflow (do this once per screenshot)").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.row_dimensions[row].height = 22
    row += 1
    workflow = [
        "1. Open claude.ai web with your connector installed.",
        "2. Start a fresh chat (no prior conversation clutter).",
        "3. Set browser zoom to ensure resulting screenshot ≥ 1000 px wide (Ctrl++ to zoom; try 100-125%).",
        "4. Paste the prompt from your use-case draft.",
        "5. When claude.ai asks 'Always allow / Deny' for the tool, click Always allow (this avoids the prompt in future captures).",
        "6. Wait for the FULL response to render. No partial captures.",
        "7. Win+Shift+S (Snipping Tool). Drag from just above the prompt to just above the message-input box.",
        "8. Save as screenshot-N.png in a dedicated folder.",
        "9. Verify dimensions are ≥ 1000 px wide (use the verify command in the Hard Requirements table above).",
        "10. Visually inspect: prompt visible? Result complete? No sidebar / chrome clutter?",
        "11. Mark row as Done in the tracker above.",
    ]
    for step in workflow:
        ws.cell(row=row, column=1, value=step)
        ws.cell(row=row, column=1).alignment = ALIGN_WRAP
        ws.cell(row=row, column=1).fill = TIP_FILL
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.row_dimensions[row].height = 28
        row += 1

    widths = [5, 26, 40, 32, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------- Tab 6: Form page-by-page ----------


def build_tab_form(wb) -> None:
    ws = wb.create_sheet("6. Form Page-by-Page")
    headers = ["Form Page", "Field", "Field Type", "Guidance", "Your Answer (draft here, paste into form)"]
    ws.append(headers)
    style_header_row(ws, 1, 5)

    SECTION = object()

    rows = [
        (SECTION, "PAGE 1 — Company Information"),
        ("Page 1", "Company/Organization Name", "Required text",
         "Your legal business name. Doesn't have to match the product name.", ""),
        ("Page 1", "Company/Organization URL", "Required text",
         "Your company's primary URL — typically the product homepage or company landing page.", ""),
        ("Page 1", "Primary Contact Name", "Required text",
         "The person reviewers will contact. Use a real name, not a role.", ""),
        ("Page 1", "Primary Contact Email", "Required text",
         "A mailbox the contact MONITORS. Reviewer correspondence happens here. Avoid info@ or similar.", ""),
        ("Page 1", "Primary Contact Role", "Optional text",
         "Job title. CTO / Founder / Engineering Lead / etc.", ""),
        ("Page 1", "Anthropic Point of Contact", "Optional text",
         "Leave blank unless you have a specific Anthropic contact.", ""),

        (SECTION, "PAGE 1 — Server Details"),
        ("Page 1", "MCP Server Name", "Required text",
         'Do NOT include "MCP" or "Server" in the name (form rule). Keep it short and brandable. '
         'If you can match your live serverInfo.title field, that\'s ideal.', ""),
        ("Page 1", "MCP Server URL type", "Required radio",
         "Universal URL = one URL serves all users (most common). Custom URLs = each user gets their own URL.", ""),
        ("Page 1", "MCP Server URL", "Required text",
         "Your live HTTPS endpoint. Must include the /mcp path (or whatever your server expects).", ""),
        ("Page 1", "Tagline", "Required text (≤ 55 chars)",
         "Short blurb appearing below your server name. Count characters carefully. Use Tab 7 template as starting point.", ""),
        ("Page 1", "MCP Server Description", "Required text (50-100 words)",
         "What your server does + key capabilities. Word count is enforced. Use Tab 7 template.", ""),
        ("Page 1", "Use Cases + Examples", "Required text (≥ 3 use cases)",
         "Each use case = 1-line description + paste-ready prompt + expected result shape. URL-bearing prompts are best.", ""),
        ("Page 1", "Connection requirements", "Required text",
         'If nothing special, write "No special requirements." Otherwise describe what the user needs '
         '(account type, geographic limits, admin seat, etc.).', ""),

        (SECTION, "PAGE 1 — Capability Classification"),
        ("Page 1", "Read/Write Capabilities", "Required radio",
         "Read Only = tools only read user data. Write Only = tools only create/modify data. "
         "Read+Write = both. If your tools transform user files (redact, blur, summarize), that's "
         "typically Read+Write.", ""),
        ("Page 1", 'Is this an "MCP App"?', "Required radio",
         "Yes if you implement interactive UI elements (ui/open-link or similar). Most MCP servers = No.", ""),
        ("Page 1", "Third-party Connections — Web access", "Multi-select",
         "Check if any tool fetches from arbitrary URLs on the open web. Includes server-side URL fetching, "
         "web scraping, search.", ""),
        ("Page 1", "Third-party Connections — AI model integration", "Multi-select",
         "Check if you send data to or receive from an AI model other than Claude (OpenAI, Gemini, Llama, etc.).", ""),
        ("Page 1", "Third-party Connections — Data retrieval", "Multi-select",
         "Check if you connect to external services that themselves retrieve data from other sources "
         "(workflow platforms, data aggregators).", ""),
        ("Page 1", "Third-party Connections — Data modification", "Multi-select",
         "Check if you can connect to external services that modify data in third-party systems.", ""),
        ("Page 1", "Data Handling — Server only accesses user-requested data", "Checkbox",
         "Check after verifying no tools fetch data the user didn't reference.", ""),
        ("Page 1", "Data Handling — No data stored beyond session", "Checkbox",
         "Check if you only store metadata (org_id, tool name, timestamp), not actual user content.", ""),
        ("Page 1", "Data Handling — Encrypted in transit", "Checkbox",
         "Check if all your endpoints use HTTPS/TLS.", ""),
        ("Page 1", "Data Handling — GDPR compliant", "Checkbox",
         "Check only if your privacy policy is genuinely GDPR-compliant. Don't check for vibes.", ""),
        ("Page 1", "Personal health data?", "Required radio",
         "Yes only if you handle medical records / lab results / health metrics. Most servers = No.", ""),
        ("Page 1", "Category", "Required radio",
         "Pick best fit: Business & Productivity / Communication / Data & Analytics / Development tools / "
         "Financial Services / Consumer Health / Health & Life Sciences / Media & Entertainment / "
         "Commerce & Shopping / Other.", ""),
        ("Page 1", "Sponsored content / ads?", "Required radio",
         "No, banner ads, sponsored ranking. Most submissions = No.", ""),

        (SECTION, "PAGE 1 — Authentication"),
        ("Page 1", "Authentication Type", "Required radio",
         "Options: No auth needed / OAuth 2.0 / Custom URL (not supported). Pick OAuth 2.0 unless truly no auth. "
         "See Tab 4 for the auth-matching playbook if your real flow is API-key-based.", ""),
        ("Page 1", "Auth Client", "Required radio",
         "Static OAuth Client = you have a pre-registered client. Dynamic OAuth Client (DCR/CIMD) = client "
         "self-registers. If you don't have a pre-registered client, pick Dynamic.", ""),
        ("Page 1", "Static Client ID (if applicable)", "Optional text",
         'If you have a pre-registered OAuth client, paste its client_id. If you picked Dynamic OAuth Client '
         'OR are using API key auth, LEAVE BLANK and explain in reviewer-instructions on page 2.', ""),
        ("Page 1", "Static Client Secret (if applicable)", "Optional text",
         'Same as above — paste the pre-registered secret only if you have one. Otherwise leave blank.', ""),
        ("Page 1", "Transport Support", "Multi-select checkboxes",
         "Streamable HTTP yes (recommended). SSE check only if your server supports SSE upgrade.", ""),

        (SECTION, "PAGE 1 — Documentation & Support"),
        ("Page 1", "MCP Server Documentation Link", "Required text",
         "Your /docs URL. Must be publicly accessible (no login).", ""),
        ("Page 1", "Privacy Policy", "Required text",
         "Your /privacy URL. Must cover data collection, use, retention, processors.", ""),
        ("Page 1", "Data Processing Agreement URL", "Optional text",
         "Leave blank unless you have a DPA published.", ""),
        ("Page 1", "Support Channel", "Required text",
         "Your /support URL or support email. Must be different from docs URL.", ""),

        (SECTION, "PAGE 2 — Test Credentials"),
        ("Page 2", "Testing Account Credentials", "Required text",
         "BARE CREDENTIALS: API key / OAuth client_id / username+password. Tell the reviewer EXACTLY what "
         "header to set (e.g., 'Authorization: Bearer <key>'). Use Tab 7 'Testing Account Credentials' template. "
         "Substitute placeholders with real values BEFORE submitting.", ""),
        ("Page 2", "Test Account Server URL (if different from main)", "Optional text",
         "Leave blank if testing happens on the same /mcp URL as production. Use only if you have a separate staging URL.", ""),
        ("Page 2", "Test Account Setup Instructions", "Required text",
         "LONG NARRATIVE: Quick start (60 seconds), 5-prompt sweep with expected results, troubleshooting, "
         "direct verification instructions. THIS IS THE HIGHEST-LEVERAGE FIELD ON THE FORM. Use Tab 7 "
         "'Test Account Setup Instructions' template.", ""),

        (SECTION, "PAGE 3 — Server Inventory"),
        ("Page 3", "List of tools in your MCP Server", "Required text",
         "Comma-separated, one line. Format: tool_name (human-readable name), tool_name (human-readable name), ... "
         "Pull names from your live tools/list response. The human-readable name should match the tool's title field.", ""),
        ("Page 3", "List of resources in your MCP Server", "Optional text",
         'Type "None" if you don\'t implement MCP Resources. (Most servers don\'t.)', ""),
        ("Page 3", "List of prompts in your MCP Server", "Optional text",
         'Type "None" if you don\'t implement MCP Prompts. (Note: MCP Prompts are reusable parameterized '
         'templates per the spec, NOT system prompts. Different concept.)', ""),
        ("Page 3", "Tool Titles & Annotations checkboxes", "2 checkboxes",
         "Check both after verifying every tool has a title field and all 4 annotations set.", ""),

        (SECTION, "PAGE 4 — Branding & Visuals"),
        ("Page 4", "Server Logo", "Required text (URL)",
         "1:1 square SVG (preferred) or PNG. Paste the public URL. Must return 200 — test with curl FIRST.", ""),
        ("Page 4", "Favicon verification", "Checkbox",
         "Verify https://www.google.com/s2/favicons?domain=YOUR-DOMAIN&sz=64 returns your favicon. "
         "Google's cache refreshes every 24-48h.", ""),
        ("Page 4", "Promotional Screenshots", "File upload (3-5 PNG, ≥ 1000 px)",
         "Upload the screenshots from Tab 5. Each ≥ 1000 px wide, PNG, shows tool output paired with prompt.", ""),
        ("Page 4", "Google Drive folder (optional)", "Optional text",
         "Optional. A Drive folder linking promo assets + matching prompts. Skip if not needed.", ""),

        (SECTION, "PAGE 4 — Launch Readiness"),
        ("Page 4", "Tested in Claude.ai (web)", "Checkbox",
         "Check only if you've actually tested. The screenshot from Tab 5 is your proof.", ""),
        ("Page 4", "Tested in Claude Desktop", "Checkbox",
         "Optional per the form. Check only if verified.", ""),
        ("Page 4", "Tested in Claude Code / Cowork", "Checkbox",
         "Optional per the form. Check only if verified.", ""),
        ("Page 4", "Server GA Date", "Required date",
         "When did your service go generally available? Pick a real date.", ""),

        (SECTION, "PAGE 5 — Skills & Plugins"),
        ("Page 5", "All fields", "All optional",
         "Form text says NOT required for MCP server submission. Leave all fields blank unless you're "
         "specifically submitting a Skill or Plugin alongside.", ""),

        (SECTION, "PAGE 6 — Submission Requirements Checklist"),
        ("Page 6", "Policy Compliance (5 checkboxes)", "5 checkboxes",
         "All 5 should be tickable after verifying: agree to Policy / no coercive automation / no financial "
         "transactions / live and ready for production / you control all API endpoints.", ""),
        ("Page 6", "Technical Requirements (6 checkboxes)", "6 checkboxes",
         "All 6 should be tickable after Tab 2 pre-flight: OAuth 2.0 implemented / tool annotations / HTTPS / "
         "CORS / IP allowlist (N/A is fine) / tested in claude.ai latest.", ""),
        ("Page 6", "Documentation Requirements (4 checkboxes)", "4 checkboxes",
         "Docs published / setup+troubleshooting / privacy policy / terms of service. All Tab 3 URLs.", ""),
        ("Page 6", "Testing Requirements (3 checkboxes)", "3 checkboxes",
         "Test account ready / credentials valid 30+ days / tools functional in target surfaces.", ""),
        ("Page 6", "Additional Information", "Optional text",
         "Use this field to surface anything the form didn't have a clean place for. Auth-mismatch explanation, "
         "design rationale, anything reviewer should know. Use Tab 7 'Additional Information' template.", ""),
        ("Page 6", "Submit button", "Action",
         "FINAL CHECK before clicking: walk back through every page; no placeholder strings <PASTE_X> survive in any field.", ""),
    ]

    row = 2
    for entry in rows:
        if entry[0] is SECTION:
            style_section_row(ws, row, 5, entry[1])
            row += 1
            continue
        page, field, ftype, guidance, _ = entry
        ws.cell(row=row, column=1, value=page)
        ws.cell(row=row, column=2, value=field)
        ws.cell(row=row, column=3, value=ftype)
        ws.cell(row=row, column=4, value=guidance)
        ws.cell(row=row, column=5, value="")
        style_data_row(ws, row, 5)
        # Highlight CRITICAL leverage fields
        if "HIGHEST-LEVERAGE" in guidance.upper() or "HIGHEST-STAKES" in guidance.upper():
            for c in range(1, 6):
                ws.cell(row=row, column=c).fill = CRITICAL_FILL
        row += 1

    widths = [11, 38, 22, 60, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, row):
        ws.row_dimensions[r].height = 48
    ws.freeze_panes = "A2"


# ---------- Tab 7: Paste-ready templates ----------


def build_tab_templates(wb) -> None:
    ws = wb.create_sheet("7. Paste-Ready Templates")

    ws["A1"] = "Paste-Ready Templates"
    ws["A1"].font = Font(bold=True, size=14, color="1E3A8A")
    ws.merge_cells("A1:B1")

    ws["A2"] = (
        "Customize the placeholders in CAPS_AND_UNDERSCORES with your product's values. Paste the result into "
        "the matching form field. Word counts and char counts are written assuming you'll keep the template "
        "structure intact."
    )
    ws["A2"].alignment = ALIGN_WRAP
    ws.fill = None
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 50

    templates = [
        ("Tagline (Form: Tagline, ≤ 55 chars)",
         "PRODUCT_CATEGORY for AI agents.\n\n"
         "(Example: 'Document and image intelligence for AI agents.' = 46 chars)\n\n"
         "Count chars carefully. Use a char counter. Aim 40-50 chars to be safe."),

        ("Description (Form: MCP Server Description, 50-100 words)",
         "PRODUCT_NAME gives AI agents N tools for PROBLEM_DOMAIN: TOOL_LIST. KEY_DESIGN_PROPERTY (one sentence). "
         "Powered by MODEL_PROVIDER, with DEPENDENCY_2 for SUB_CAPABILITY_1 and DEPENDENCY_3 for SUB_CAPABILITY_2. "
         "TENANCY_MODEL with org-scoped quotas, rate limits, and atomic metering. Free tier available; "
         "AUTH_METHODS supported.\n\n"
         "Word count check: paste into Word/Docs/an online counter. Stay 60-90 for safety."),

        ("Use Cases (Form: Use Cases + Examples, ≥ 3)",
         "1. USE_CASE_1_NAME\n"
         '   "Use PRODUCT_NAME to ACTION at https://YOUR-DOMAIN/samples/sample1.ext — EXPECTED_OUTPUT_DESCRIPTION."\n'
         "   Returns: { field1, field2, field3 }\n\n"
         "2. USE_CASE_2_NAME\n"
         '   "Use PRODUCT_NAME to ACTION at https://YOUR-DOMAIN/samples/sample2.ext — EXPECTED_OUTPUT_DESCRIPTION."\n'
         "   Returns: { field1, field2, field3 }\n\n"
         "3. USE_CASE_3_NAME\n"
         '   "Use PRODUCT_NAME to ACTION at https://YOUR-DOMAIN/samples/sample3.ext — EXPECTED_OUTPUT_DESCRIPTION."\n'
         "   Returns: { field1, field2, field3 }\n\n"
         "(Optional 4-5: copy the pattern. URL-bearing prompts work best because they avoid the chat-host "
         "base64 stall on file uploads.)"),

        ("Connection Requirements (Form: Connection requirements)",
         "No special requirements. Add the connector at https://YOUR-DOMAIN/mcp and paste the provided API key "
         "as a Bearer token. No admin seat needed, no custom URL, no geographic restriction. Sample-file URLs "
         "work from anywhere."),

        ("Testing Account Credentials (Form: Testing Account Credentials)",
         "Authentication: AUTH_METHOD (e.g., Bearer API key, preferred for review)\n\n"
         "Credential: <PASTE_AT_SUBMISSION_TIME>\n"
         "Header to send: Authorization: Bearer <credential>\n\n"
         "This credential is bound to a free-tier org/account with QUOTA_DETAILS. Issued for the test account "
         "TEST_EMAIL.\n\n"
         "Test account (for dashboard access if needed):\n"
         "  Email: TEST_EMAIL\n"
         "  Password: <PASTE_AT_SUBMISSION_TIME>\n"
         "  Dashboard: https://YOUR-DOMAIN/dashboard\n"
         "  2FA: not required.\n\n"
         "Note (if auth-mismatch applies): the 'Static Client ID / Secret' fields on page 1 were left blank "
         "intentionally — PRODUCT_NAME does not have a pre-registered static OAuth client. The form's auth "
         "options didn't cleanly match our setup. See the test-account-setup-instructions field for full details."),

        ("Test Account Setup Instructions (Form: Test Account Setup Instructions — THE BIG ONE)",
         "## Quick start (60 seconds)\n\n"
         "1. In claude.ai, open Settings → Connectors → Add custom connector.\n"
         "2. Paste this URL: https://YOUR-DOMAIN/mcp\n"
         "3. Open Advanced Settings. Configure auth as **Bearer token** and paste:\n\n"
         "      <PASTE_API_KEY_AT_SUBMISSION_TIME>\n\n"
         "   (Credential bound to a free-tier org for the test account TEST_EMAIL.)\n\n"
         "4. Click Add. All N tools should appear immediately: TOOL_LIST.\n\n"
         "No sign-in flow, no popup. Sample files are hosted at public URLs (see below) — no upload needed.\n\n"
         "## N-prompt sweep — paste each prompt into a fresh chat\n\n"
         "Each prompt names a public PRODUCT_NAME-hosted sample-file URL. Total quota burn: N of N free-tier calls.\n\n"
         "1. TOOL_1_NAME\n"
         '   Prompt: "Use PRODUCT_NAME to ACTION at https://YOUR-DOMAIN/samples/sample1.ext — EXPECTED."\n'
         "   Expect: DESCRIPTION_OF_VISUAL_RESULT.\n\n"
         "2. TOOL_2_NAME\n"
         '   Prompt: "Use PRODUCT_NAME to ACTION at https://YOUR-DOMAIN/samples/sample2.ext."\n'
         "   Expect: DESCRIPTION_OF_VISUAL_RESULT.\n\n"
         "(... continue for all your sweep tools ...)\n\n"
         "## Direct verification (no claude.ai required)\n\n"
         "  npx @modelcontextprotocol/inspector\n"
         "  → Transport: Streamable HTTP\n"
         "  → URL: https://YOUR-DOMAIN/mcp\n"
         "  → Auth: Bearer + the credential above\n"
         "  → Click Connect.\n"
         "All N tools render with full schemas and annotations.\n\n"
         "## Error envelope (for reference)\n\n"
         "- ERROR_CODE_1: DESCRIPTION (action: HOW_TO_RECOVER)\n"
         "- ERROR_CODE_2: DESCRIPTION\n"
         "- ERROR_CODE_3: DESCRIPTION\n\n"
         "Tool-internal failures come back as isError: true in the result envelope (not a JSON-RPC error) "
         "so Claude can recover. Quota refunded on exception.\n\n"
         "## Documentation + contact\n\n"
         "- Tool docs: https://YOUR-DOMAIN/docs\n"
         "- Privacy: https://YOUR-DOMAIN/privacy\n"
         "- Support during review: CONTACT_EMAIL\n"
         "- Security: https://YOUR-DOMAIN/security"),

        ("Server Inventory — Tools (Form: List of tools)",
         "Comma-separated, one line, format `tool_name (Human-Readable Name)`:\n\n"
         "tool1_name (Tool 1 Human Name), tool2_name (Tool 2 Human Name), tool3_name (Tool 3 Human Name)\n\n"
         "Pull human-readable names from your live tools/list response's 'title' field. They should match what "
         "appears in claude.ai's Customize → Connectors page after installing your connector."),

        ("Additional Information (Form: Additional Information on Page 6)",
         "A few notes the form didn't have a clean place for:\n\n"
         "1. KEY_DESIGN_PROPERTY. EXPLAIN_WHY_AND_WHY_REVIEWER_SHOULD_CARE.\n\n"
         "2. Auth choice rationale. The form's Auth section asked for OAuth 2.0 + AUTH_CLIENT_TYPE; the Static "
         "Client ID/Secret fields were left blank intentionally because PRODUCT_NAME uses AUTH_REAL_FLOW. See "
         "the test-account-setup-instructions field for operational details.\n\n"
         "3. ANY_OTHER_NOTES (e.g., 'tool count is N, not M because we added X mid-build', 'support hours are X', etc.)"),

        ("Reviewer-facing rationale (use for any 'tell us more' field)",
         "PRODUCT_NAME is a PRODUCT_CATEGORY MCP server. The N tools are SCOPED to PROBLEM_DOMAIN — the "
         "surface is intentionally narrow because OPTIMIZATION_GOAL.\n\n"
         "DOWNSTREAM_DEPENDENCY (e.g., AI model provider) is a downstream service we call server-to-server "
         "with our own API key (analogous to how SaaS apps call AWS/Stripe). All tool inputs are user-supplied; "
         "the server doesn't fetch data on behalf of the user without explicit invocation.\n\n"
         "AUTH is implemented via AUTH_PROVIDER. TENANCY_MODEL is enforced at the data layer; tests in "
         "PATH_TO_TESTS verify no cross-tenant data leakage on any API path."),
    ]

    row = 4
    for header, content in templates:
        # Section header for the template
        ws.cell(row=row, column=1, value=header).font = Font(bold=True, size=11, color="FFFFFF")
        ws.cell(row=row, column=1).fill = HEADER_FILL
        ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.row_dimensions[row].height = 24
        row += 1
        # Template body
        ws.cell(row=row, column=1, value=content)
        ws.cell(row=row, column=1).alignment = ALIGN_WRAP
        ws.cell(row=row, column=1).border = BORDER
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        # Estimate row height by line count
        line_count = content.count("\n") + 1
        ws.row_dimensions[row].height = max(50, min(line_count * 15, 480))
        row += 1
        # Spacer
        row += 1

    widths = [90, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------- Tab 8: Post-submission ----------


def build_tab_post_submission(wb) -> None:
    ws = wb.create_sheet("8. Post-Submission")

    ws["A1"] = "Post-Submission Tracker"
    ws["A1"].font = Font(bold=True, size=14, color="1E3A8A")
    ws.merge_cells("A1:E1")

    ws["A2"] = (
        "Don't disappear after clicking Submit. Reviewers may have questions. Production must stay stable. "
        "Track everything below until the listing is published."
    )
    ws["A2"].alignment = ALIGN_WRAP
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 40

    # Immediate actions
    ws["A4"] = "Within the first hour of submitting"
    ws["A4"].font = SECTION_FONT
    ws["A4"].fill = SECTION_FILL
    ws.merge_cells("A4:E4")
    ws.row_dimensions[4].height = 22

    headers = ["#", "Action", "Why", "Done?", "Notes"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=5, column=i, value=h)
    style_header_row(ws, 5, 5)

    immediate = [
        ("Save a screenshot of the Google Forms confirmation page",
         "Your proof of submission and timestamp."),
        ("Save the confirmation email if Google Forms sends one",
         "The form-submitter email receives it. Forward to the primary contact email."),
        ("Re-curl EVERY URL you submitted in the form",
         "Catch any URLs that broke between draft and submission. Especially the logo. "
         "If you find any 404s, fix them IMMEDIATELY."),
        ("Verify test credentials still work via tools/list",
         "Catch credential expiry / revocation."),
        ("Verify test-account quota is still ≥ 15 calls",
         "Don't let post-submission testing exhaust it before the reviewer can run their sweep."),
        ("Add a calendar reminder for +7 days, +14 days, +28 days",
         "Typical review window. If you hear nothing by +28 days, follow up with a polite check-in."),
        ("Disable any code-side automations that might rotate the API key",
         "Reviewer may test 2-4 weeks out. Don't break their auth mid-review."),
        ("Commit + push any post-submission fixes (e.g., logo URL fix) on a CLEAN BRANCH separate from MCP-surface changes",
         "Don't drift the deployed state from what was submitted. Static-asset fixes are fine; tool-surface changes are not."),
    ]
    row = 6
    for i, (action, why) in enumerate(immediate, 1):
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=action)
        ws.cell(row=row, column=3, value=why)
        ws.cell(row=row, column=4, value="")
        ws.cell(row=row, column=5, value="")
        style_data_row(ws, row, 5)
        ws.row_dimensions[row].height = 45
        row += 1
    add_yesno_validation(ws, f"D6:D{row-1}")

    # Ongoing during review
    row += 1
    ws.cell(row=row, column=1, value="Ongoing during review (next 2-4 weeks)").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.row_dimensions[row].height = 22
    row += 1

    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, 5)
    row += 1
    ongoing_start = row

    ongoing = [
        ("Don't push code changes to the MCP surface",
         "Drift between submission state and deployed state creates avoidable failure modes."),
        ("Monitor primary contact email daily",
         "Reviewer correspondence may arrive at any time. Slow response = slow review."),
        ("Don't promote / market the directory listing until approved",
         "Pre-announcing creates expectations; if rejected, you have to walk it back."),
        ("Keep production endpoint healthy",
         "If your /mcp endpoint goes down during review, you may get bounced."),
        ("Keep documentation links healthy",
         "Same. Reviewer may revisit /docs / /privacy / /support during review."),
        ("If you discover a bug, decide carefully whether to deploy a fix",
         "Tool-surface bugs may justify a redeploy. Cosmetic / non-blocking bugs can wait until post-review. "
         "Document your decision."),
        ("If reviewer requests changes, respond within 48 hours",
         "Fast iteration = faster approval."),
    ]
    for i, (action, why) in enumerate(ongoing, 1):
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=action)
        ws.cell(row=row, column=3, value=why)
        ws.cell(row=row, column=4, value="")
        ws.cell(row=row, column=5, value="")
        style_data_row(ws, row, 5)
        ws.row_dimensions[row].height = 45
        row += 1
    add_yesno_validation(ws, f"D{ongoing_start}:D{row-1}")

    # After approval
    row += 1
    ws.cell(row=row, column=1, value="After approval").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.row_dimensions[row].height = 22
    row += 1

    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, 5)
    row += 1
    after_start = row

    after = [
        ("Verify the listing renders correctly in the directory",
         "Logo / tagline / description / screenshots / links all visible and correct."),
        ("Test the install flow as a brand-new user",
         "Fresh account, no prior history. Does it work end-to-end?"),
        ("Announce / market the listing",
         "Now you can. Blog post, social, email list, etc."),
        ("Stand up monitoring on the prod endpoint",
         "Uptime / response time / error rate. Listing means real users will hit it."),
        ("Plan for v1.1 improvements",
         "Anything you punted to ship the v1 submission. SSE streaming, more tools, OAuth client publish, etc."),
        ("Update the listing if your tool surface evolves",
         "Anthropic Software Directory Policy requires keeping the listing in sync with what's deployed. "
         "Add/remove tools = update listing."),
    ]
    for i, (action, why) in enumerate(after, 1):
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=action)
        ws.cell(row=row, column=3, value=why)
        ws.cell(row=row, column=4, value="")
        ws.cell(row=row, column=5, value="")
        style_data_row(ws, row, 5)
        ws.row_dimensions[row].height = 45
        row += 1
    add_yesno_validation(ws, f"D{after_start}:D{row-1}")

    widths = [5, 50, 55, 9, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Final cell: who-to-contact
    row += 1
    ws.cell(row=row, column=1, value="If review takes longer than 4 weeks").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.row_dimensions[row].height = 22
    row += 1
    contact_text = (
        "Send a polite follow-up to the email Anthropic gave you in the confirmation, or to "
        "mcp-review@anthropic.com (the published reviewer mailbox). Include: your submission date, your product "
        "name, and a brief 'just checking in' note. Don't be pushy; reviewers are busy."
    )
    ws.cell(row=row, column=1, value=contact_text)
    ws.cell(row=row, column=1).alignment = ALIGN_WRAP
    ws.cell(row=row, column=1).fill = TIP_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.row_dimensions[row].height = 65


# ---------- Main ----------


def main() -> None:
    wb = openpyxl.Workbook()
    build_tab_readme(wb)
    build_tab_preflight(wb)
    build_tab_public_site(wb)
    build_tab_auth(wb)
    build_tab_screenshots(wb)
    build_tab_form(wb)
    build_tab_templates(wb)
    build_tab_post_submission(wb)

    downloads = Path(os.path.expanduser("~")) / "Downloads"
    out_path = downloads / "MCP_Directory_Submission_Checklist.xlsx"
    wb.save(out_path)
    print(f"Wrote: {out_path}")
    print(f"Tabs: {wb.sheetnames}")


if __name__ == "__main__":
    main()
