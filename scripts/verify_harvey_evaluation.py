"""Verify six live tools with a temporary key for the dedicated evaluation org.

This validates organization-level service access, not unverified users' OAuth
sign-ins. Uses synthetic documents and the existing public reviewer images.
"""
from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import secrets
import time
from datetime import datetime, timezone

import requests
from docx import Document
from dotenv import dotenv_values
from PIL import Image
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from db.models import ApiKey, Org, Quota
from scripts.provision_harvey_evaluation import BASE_URL, ORG_NAME, ROOT, STATE_PATH, emit, save


def main():
    os.environ["REQUESTS_CA_BUNDLE"] = str(ROOT / ".build_notes/harvey/windows-ca.pem")
    cfg = dotenv_values(ROOT / ".env")
    state = json.loads(STATE_PATH.read_text(encoding="utf-8"))
    db_url = cfg["DATABASE_URL"].replace("postgresql://", "postgresql+psycopg://", 1)
    db_url = db_url.replace("postgres://", "postgresql+psycopg://", 1)
    engine = create_engine(db_url, connect_args={"connect_timeout": 10})
    started = datetime.now(timezone.utc)
    with Session(engine) as session:
        org = session.get(Org, state["synzo_org_id"])
        if not org or org.name != ORG_NAME or org.workos_org_id != state["workos_org_id"]:
            raise RuntimeError("Evaluation organization did not match saved state")
        for previous in session.scalars(select(ApiKey).where(
            ApiKey.org_id == org.id, ApiKey.name == "Harvey provisioning validation",
            ApiKey.revoked_at.is_(None))):
            previous.revoked_at = started
        raw_key = "sk_synzo_" + secrets.token_urlsafe(32)
        record = ApiKey(org_id=org.id, name="Harvey provisioning validation",
                        key_hash=hashlib.sha256(raw_key.encode()).hexdigest(), prefix=raw_key[:16])
        session.add(record)
        session.commit()
        key_id = record.id
    state["validation_api_key_id"] = key_id
    save(state)
    http = requests.Session()
    headers = {"Authorization": "Bearer " + raw_key,
               "Accept": "application/json, text/event-stream"}
    results = []

    def call(name, arguments, required):
        start = time.monotonic()
        try:
            response = http.post(BASE_URL + "/mcp", headers=headers, json={
                "jsonrpc": "2.0", "id": len(results) + 1, "method": "tools/call",
                "params": {"name": name, "arguments": arguments},
            }, timeout=90)
            payload = response.json()
            result = payload.get("result", {})
            content = result.get("structuredContent", {})
            success = response.ok and not result.get("isError", True) and all(content.get(k) for k in required)
            item = {"tool": name, "success": bool(success), "http_status": response.status_code,
                    "seconds": round(time.monotonic() - start, 2)}
            if not success:
                item["rpc_error"] = payload.get("error", {}).get("code")
                item["tool_error"] = next((c.get("text", "")[:500] for c in result.get("content", [])
                                           if c.get("type") == "text"), "")
            results.append(item)
            emit(**item)
            return content if success else None
        except (requests.RequestException, ValueError) as error:
            item = {"tool": name, "success": False, "error_type": type(error).__name__,
                    "seconds": round(time.monotonic() - start, 2)}
            results.append(item)
            emit(**item)
            return None

    try:
        doc = Document()
        doc.add_heading("Synthetic connector evaluation memorandum", 0)
        doc.add_paragraph("This fictional document is for testing Synzo's Harvey connector. "
                          "It contains no real client or matter information.")
        doc.add_paragraph("The sample agreement runs for twelve months. Either party may terminate "
                          "with thirty days' written notice. The fee is 100 Canadian dollars per month.")
        doc.add_paragraph("Synthetic contact: Jane Doe, jane.doe@example.com, telephone 416-555-0100.")
        buffer = io.BytesIO()
        doc.save(buffer)
        upload = call("upload_file", {"filename": "harvey-synthetic-evaluation.docx",
            "content_base64": base64.b64encode(buffer.getvalue()).decode()}, ["content_url"])
        if upload:
            arguments = {"filename": "harvey-synthetic-evaluation.docx", "content_url": upload["content_url"]}
            call("summarize_document", arguments, ["summary"])
            call("translate_document", {**arguments, "target_language": "French"}, ["translated_text"])
            redacted = call("redact_pii", arguments, ["result_url"])
            if redacted:
                output = http.get(redacted["result_url"], timeout=30)
                output.raise_for_status()
                processed = Document(io.BytesIO(output.content))
                text = "\n".join(p.text for p in processed.paragraphs)
                results[-1]["download_verified"] = True
                results[-1]["synthetic_email_removed"] = "jane.doe@example.com" not in text
                results[-1]["success"] = results[-1]["success"] and results[-1]["synthetic_email_removed"]
                emit(tool="redact_pii", download_verified=True,
                     synthetic_email_removed=results[-1]["synthetic_email_removed"])
        else:
            for name in ("summarize_document", "translate_document", "redact_pii"):
                results.append({"tool": name, "success": False, "reason": "upload_failed"})
        call("analyze_image", {"filename": "analyze-sample.jpg",
            "content_url": BASE_URL + "/static/reviewer-samples/analyze-sample.jpg"}, ["analysis"])
        faces = call("detect_faces", {"filename": "detect-faces-sample.jpg", "mode": "redact",
            "content_url": BASE_URL + "/static/reviewer-samples/detect-faces-sample.jpg"}, ["result_url"])
        if faces:
            output = http.get(faces["result_url"], timeout=30)
            output.raise_for_status()
            picture = Image.open(io.BytesIO(output.content))
            picture.verify()
            results[-1]["download_verified"] = picture.format == "PNG"
            results[-1]["success"] = results[-1]["success"] and picture.format == "PNG"
            emit(tool="detect_faces", png_download_verified=picture.format == "PNG")
    finally:
        with Session(engine) as session:
            record = session.get(ApiKey, key_id)
            record.revoked_at = datetime.now(timezone.utc)
            session.commit()
            now = datetime.now(timezone.utc)
            quota = session.scalar(select(Quota).where(Quota.org_id == state["synzo_org_id"],
                                                        Quota.period_start <= now, Quota.period_end > now))
            remaining = quota.calls_remaining
        state["service_verification"] = {"started_at": started.isoformat(),
            "completed_at": datetime.now(timezone.utc).isoformat(), "results": results,
            "temporary_api_key_revoked": True, "calls_remaining": remaining,
            "all_six_tools_passed": len(results) == 6 and all(r["success"] for r in results)}
        save(state)
        emit(temporary_api_key_revoked=True, calls_remaining=remaining,
             all_six_tools_passed=state["service_verification"]["all_six_tools_passed"])


if __name__ == "__main__":
    main()
